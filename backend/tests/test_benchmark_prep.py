from __future__ import annotations

import csv
import json
import sqlite3
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from io import StringIO
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT_ROOT / "ml_pipeline" / "scripts"))

import openpyxl  # noqa: E402

from build_language_db import (  # type: ignore  # noqa: E402
    EVAL_HOLDOUT_SPLIT,
    apply_corpus_checks,
    connect,
    create_schema,
    default_eval_paths,
    ingest_parallel,
    load_eval_items,
)
from eval_set import main as eval_main, read_sources, source_overlaps  # type: ignore  # noqa: E402
from export_datasets import export  # type: ignore  # noqa: E402
from worksheets import main as worksheets_main  # type: ignore  # noqa: E402

# English sources are test fixtures; Zomi-side strings are placeholders, never real language data.


def write_sources(path: Path, rows: list[tuple[str, str, str]]) -> Path:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["id", "domain", "english"])
        writer.writerows(rows)
    return path


def write_zomi_sources(path: Path, rows: list[tuple[str, str, str]]) -> Path:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["id", "domain", "zomi"])
        writer.writerows(rows)
    return path


class SourcesTests(unittest.TestCase):
    def setUp(self) -> None:
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.dir = Path(tmp.name)

    def test_sources_become_english_only_items(self) -> None:
        items, errors = read_sources(write_sources(self.dir / "s.csv", [("en-0001", "health", "Take two tablets a day.")]))
        self.assertEqual(errors, [])
        self.assertEqual(
            {k: items[0][k] for k in ("id", "direction", "domain", "source_text", "references")},
            {"id": "en-0001", "direction": "en-zomi", "domain": "health",
             "source_text": "Take two tablets a day.", "references": []},
        )

    def test_zomi_sources_become_zomi_to_english_items(self) -> None:
        items, errors = read_sources(write_zomi_sources(self.dir / "z.csv", [("zomi-en-0001", "education", "ZOMI_PLACEHOLDER_A")]))
        self.assertEqual(errors, [])
        self.assertEqual((items[0]["direction"], items[0]["source_text"]), ("zomi-en", "ZOMI_PLACEHOLDER_A"))

    def test_a_sources_list_holds_one_language(self) -> None:
        both = self.dir / "both.csv"
        both.write_text("id,domain,english,zomi\nx,news,One.,ZOMI_PLACEHOLDER\n", encoding="utf-8")
        self.assertIn("one language", read_sources(both)[1][0])
        neither = self.dir / "neither.csv"
        neither.write_text("id,domain,text\nx,news,One.\n", encoding="utf-8")
        self.assertIn("missing columns english or zomi", read_sources(neither)[1][0])
        _, errors = read_sources(write_zomi_sources(self.dir / "blank.csv", [("zomi-en-0001", "news", "")]))
        self.assertIn("id and zomi are required", errors[0])

    def test_bad_sources_are_reported(self) -> None:
        rows = [("en-0001", "health", "One."), ("en-0001", "health", "Two."), ("en-0003", "sport", "Three."),
                ("en-0004", "news", "")]
        _, errors = read_sources(write_sources(self.dir / "s.csv", rows))
        self.assertEqual(len(errors), 3)
        (self.dir / "bad.csv").write_text("id,text\n1,x\n", encoding="utf-8")
        _, errors = read_sources(self.dir / "bad.csv")
        self.assertIn("missing columns", errors[0])

    def test_overlaps_are_found_exactly_and_as_number_variants(self) -> None:
        items, _ = read_sources(write_sources(self.dir / "s.csv", [
            ("a", "informal", "I have 3 cows."), ("b", "informal", "Where is the market?"), ("c", "news", "Fresh text."),
        ]))
        hits = source_overlaps(items, ["I have 7 cows.", "where is the market"])
        self.assertEqual(hits, {"a": {"variant"}, "b": {"exact", "variant"}})

    def test_default_evaluation_files_include_drafts_and_source_lists(self) -> None:
        folder = self.dir / "human"
        folder.mkdir()
        for name in ("en-zomi-v1.jsonl", "en-zomi-v1.draft.jsonl", "en-zomi-v1.sources.csv",
                     "en-zomi-v1.manifest.json", "notes.txt"):
            (folder / name).write_text("", encoding="utf-8")
        self.assertEqual([p.name for p in default_eval_paths(folder)],
                         ["en-zomi-v1.draft.jsonl", "en-zomi-v1.jsonl", "en-zomi-v1.sources.csv"])
        self.assertEqual(default_eval_paths(self.dir / "missing"), [])


class SpreadsheetSafetyTests(unittest.TestCase):
    def setUp(self) -> None:
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.dir = Path(tmp.name)

    def test_excel_csv_with_byte_order_mark_is_read(self) -> None:
        path = self.dir / "s.csv"
        path.write_bytes("\ufeffid,domain,english\nen-0001,news,The market reopened.\n".encode("utf-8"))
        items, errors = read_sources(path)
        self.assertEqual((errors, items[0]["id"]), ([], "en-0001"))

    def test_short_rows_are_reported_not_crashing(self) -> None:
        path = self.dir / "s.csv"
        path.write_text("id,domain,english\nen-0001,news\n", encoding="utf-8")
        _, errors = read_sources(path)
        self.assertIn("id and english are required", errors[0])

    def test_sheet_cells_starting_with_equals_stay_text(self) -> None:
        sources = write_sources(self.dir / "s.csv", [("en-0001", "news", '=HYPERLINK("http://example.com","x")')])
        sheet = self.dir / "t.xlsx"
        with redirect_stdout(StringIO()):
            worksheets_main(["testset-sheet", str(sources), "--out", str(sheet)])
        cell = openpyxl.load_workbook(sheet)["Translate"]["C2"]
        self.assertEqual((cell.data_type, cell.value), ("s", '=HYPERLINK("http://example.com","x")'))

        review = self.dir / "review.csv"
        review.write_text("sample_id,english,zomi,meaning,fluency,correction,notes\nR00001,Hi.,=1+1,,,,\n",
                          encoding="utf-8")
        with redirect_stdout(StringIO()):
            worksheets_main(["ratings-sheet", str(review), "--out", str(self.dir / "r.xlsx")])
        cell = openpyxl.load_workbook(self.dir / "r.xlsx")["Rate"]["C2"]
        self.assertEqual((cell.data_type, cell.value), ("s", "=1+1"))


class ExclusionTests(unittest.TestCase):
    def setUp(self) -> None:
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.dir = Path(tmp.name)
        corpus = self.dir / "corpus.csv"
        with corpus.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(["index", "en", "zomi", "model", "worker"])
            for n, english in enumerate(["The clinic opens at 9 o'clock.", "Children walk to school.", "It rained all week."]):
                writer.writerow([n, english, f"ZOMI_PLACEHOLDER_{n}", "model-a", "w"])
        self.database = self.dir / "language.sqlite"
        con = connect(self.database)
        create_schema(con)
        with redirect_stdout(StringIO()):
            ingest_parallel(con, corpus, "2026-10-25T00:00:00+00:00")
        self.sources = write_sources(self.dir / "en-zomi-v1.sources.csv", [
            ("en-0001", "health", "The clinic opens at 8 o'clock."),  # a number variant of a corpus sentence
            ("en-0002", "education", "Children walk to school."),      # an exact copy
        ])
        apply_corpus_checks(con, load_eval_items([self.sources]))
        con.close()

    def test_build_moves_source_matches_and_their_variants_to_the_holdout(self) -> None:
        con = sqlite3.connect(self.database)
        splits = dict(con.execute("SELECT normalized_source_text, dataset_split FROM parallel_sentences"))
        con.close()
        self.assertEqual(splits["The clinic opens at 9 o'clock."], EVAL_HOLDOUT_SPLIT)
        self.assertEqual(splits["Children walk to school."], EVAL_HOLDOUT_SPLIT)
        self.assertNotEqual(splits["It rained all week."], EVAL_HOLDOUT_SPLIT)

    def test_export_never_writes_a_source_match(self) -> None:
        con = sqlite3.connect(self.database)
        con.execute("UPDATE parallel_sentences SET dataset_split = 'train'")  # even if the build was run without them
        result = export(
            con, self.dir / "out", rights={"parallel_corpus": "cleared"}, allow_uncleared=False, bible=[],
            phrases=[], eval_items=load_eval_items([self.sources]), summary=None, min_acceptable=0.8, min_score=0.0,
        )
        con.close()
        exported = [json.loads(line)["english"] for line in (self.dir / "out" / "train.jsonl").read_text().splitlines()]
        self.assertEqual(exported, ["It rained all week."])
        self.assertEqual(result["skipped"]["evaluation_match"], 2)

    def test_check_sources_reports_corpus_and_bible_overlaps(self) -> None:
        bible = self.dir / "bible.sqlite"
        con = sqlite3.connect(bible)
        con.executescript(
            "CREATE TABLE versions (id TEXT, language_iso TEXT);"
            "CREATE TABLE verses (version_id TEXT, text TEXT);"
            "INSERT INTO versions VALUES ('web', 'eng'), ('tdb', 'zom');"
            "INSERT INTO verses VALUES ('web', 'Jesus wept.'), ('tdb', 'ZOMI_PLACEHOLDER_VERSE');"
        )
        con.commit()
        con.close()
        sources = write_sources(self.dir / "check.csv", [
            ("en-0001", "religion", "Jesus wept."), ("en-0002", "education", "Children walk to school."),
            ("en-0003", "news", "The bridge reopened on Monday."),
        ])
        out = StringIO()
        with redirect_stdout(out):
            code = eval_main(["check-sources", str(sources), "--database", str(self.database),
                              "--bible-versions", str(bible)])
        self.assertEqual(code, 1)
        self.assertIn("en-0001 (Bible exact, Bible variant)", out.getvalue())
        self.assertIn("en-0002 (corpus exact, corpus variant)", out.getvalue())
        self.assertNotIn("en-0003", out.getvalue().split("source sentences overlap")[0])


class ZomiSourceTests(unittest.TestCase):
    """Zomi -> English sources: kept out of training and checked against the Zomi side."""

    def setUp(self) -> None:
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.dir = Path(tmp.name)
        corpus = self.dir / "corpus.csv"
        with corpus.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(["index", "en", "zomi", "model", "worker"])
            for n, english in enumerate(["The clinic opens at 9 o'clock.", "Children walk to school.", "It rained all week."]):
                writer.writerow([n, english, f"ZOMI PLACEHOLDER SENTENCE {n}", "model-a", "w"])
        self.database = self.dir / "language.sqlite"
        con = connect(self.database)
        create_schema(con)
        with redirect_stdout(StringIO()):
            ingest_parallel(con, corpus, "2026-10-25T00:00:00+00:00")
        con.close()
        self.sources = write_zomi_sources(self.dir / "zomi-en-v1.sources.csv", [
            ("zomi-en-0001", "education", "ZOMI PLACEHOLDER SENTENCE 1"),  # the Zomi side of a corpus pair
            ("zomi-en-0002", "informal", "ZOMI PLACEHOLDER ONLY IN THE BOOK"),
        ])

    def test_build_keeps_pairs_with_a_zomi_source_out_of_training(self) -> None:
        con = sqlite3.connect(self.database)
        apply_corpus_checks(con, load_eval_items([self.sources]))
        splits = dict(con.execute("SELECT normalized_source_text, dataset_split FROM parallel_sentences"))
        con.close()
        self.assertEqual(splits["Children walk to school."], EVAL_HOLDOUT_SPLIT)
        self.assertNotEqual(splits["It rained all week."], EVAL_HOLDOUT_SPLIT)
        self.assertIn(self.sources, default_eval_paths(self.dir))

    def test_check_sources_compares_zomi_with_the_zomi_side_and_zomi_bibles(self) -> None:
        bible = self.dir / "bible.sqlite"
        con = sqlite3.connect(bible)
        con.executescript(
            "CREATE TABLE versions (id TEXT, language_iso TEXT);"
            "CREATE TABLE verses (version_id TEXT, text TEXT);"
            "INSERT INTO versions VALUES ('web', 'eng'), ('tdb', 'zom');"
            "INSERT INTO verses VALUES ('web', 'ZOMI PLACEHOLDER ONLY IN THE BOOK'), ('tdb', 'ZOMI PLACEHOLDER VERSE');"
        )
        con.commit()
        con.close()
        sources = write_zomi_sources(self.dir / "check.csv", [
            ("zomi-en-0001", "education", "ZOMI PLACEHOLDER SENTENCE 1"),
            ("zomi-en-0002", "religion", "ZOMI PLACEHOLDER VERSE"),
            ("zomi-en-0003", "informal", "ZOMI PLACEHOLDER ONLY IN THE BOOK"),  # only in an English Bible: no match
        ])
        out = StringIO()
        with redirect_stdout(out):
            code = eval_main(["check-sources", str(sources), "--database", str(self.database), "--bible-versions", str(bible)])
        report = out.getvalue()
        self.assertEqual(code, 1)
        self.assertIn("zomi-en-0001 (corpus exact, corpus variant)", report)
        self.assertIn("zomi-en-0002 (Bible exact, Bible variant)", report)
        self.assertNotIn("zomi-en-0003", report.split("source sentences overlap")[0])

    def test_zomi_sheets_round_trip_into_a_valid_zomi_to_english_set(self) -> None:
        sheet, draft = self.dir / "translate.xlsx", self.dir / "draft.jsonl"
        review, final = self.dir / "review.xlsx", self.dir / "zomi-en-v1.jsonl"
        with redirect_stdout(StringIO()):
            worksheets_main(["testset-sheet", str(self.sources), "--out", str(sheet)])
            book = openpyxl.load_workbook(sheet)
            translate = book["Translate"]
            self.assertEqual([c.value for c in translate[1]], ["ID", "Topic", "Zomi", "English translation", "Notes"])
            self.assertEqual(book["Instructions"]["A1"].value, "Siamsil test set: translate Zomi into English")
            translate["D2"] = "The children are going to school."
            translate["D3"] = "It is raining."
            book.save(sheet)
            worksheets_main(["testset-import", str(sheet), "--translator", "T01", "--origin", "primer", "--out", str(draft)])
            items = [json.loads(line) for line in draft.read_text().splitlines()]
            self.assertEqual({i["direction"] for i in items}, {"zomi-en"})
            self.assertEqual(items[0]["references"], ["The children are going to school."])

            worksheets_main(["review-sheet", str(draft), "--out", str(review)])
            book = openpyxl.load_workbook(review)
            sheet_rows = book["Review"]
            self.assertEqual(sheet_rows["F1"].value, "Corrected English")
            sheet_rows["E2"], sheet_rows["E3"], sheet_rows["F3"] = "Accept", "Correct", "It is raining hard."
            book.save(review)
            worksheets_main(["review-merge", str(draft), str(review), "--reviewer", "R01", "--out", str(final)])
        merged = [json.loads(line) for line in final.read_text().splitlines()]
        self.assertEqual([(i["direction"], i["review_outcome"], i["references"]) for i in merged], [
            ("zomi-en", "accepted", ["The children are going to school."]),
            ("zomi-en", "corrected", ["It is raining hard."]),
        ])


class TranslatorSheetTests(unittest.TestCase):
    def test_sheet_round_trips_through_testset_import(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            sources = write_sources(folder / "s.csv", [("en-0001", "health", "Rest and drink water."),
                                                      ("en-0002", "news", "The road is open again.")])
            sheet, draft = folder / "translate.xlsx", folder / "draft.jsonl"
            with redirect_stdout(StringIO()):
                worksheets_main(["testset-sheet", str(sources), "--out", str(sheet)])
                book = openpyxl.load_workbook(sheet)
                book["Translate"]["D2"] = "ZOMI_PLACEHOLDER_TRANSLATION"
                book.save(sheet)
                worksheets_main(["testset-import", str(sheet), "--translator", "T01", "--origin", "test", "--out", str(draft)])
            items = [json.loads(line) for line in draft.read_text().splitlines()]
            self.assertEqual([(i["id"], i["domain"], i["source_text"]) for i in items],
                             [("en-0001", "health", "Rest and drink water.")])


if __name__ == "__main__":
    unittest.main()
