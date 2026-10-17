from __future__ import annotations

import csv
import json
import sqlite3
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from io import StringIO
from pathlib import Path
from unittest.mock import patch

PROJECT_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT_ROOT / "ml_pipeline" / "scripts"))

import openpyxl  # noqa: E402

from build_language_db import apply_corpus_checks, connect, create_schema, ingest_parallel  # type: ignore  # noqa: E402
from export_datasets import admitted_sources, export, load_workbook_sources, main, silver_models  # type: ignore  # noqa: E402
from import_usx_bibles import main as import_bibles  # type: ignore  # noqa: E402

# Zomi-side strings are placeholders, never real language data.

ALL_NEEDS_REVIEW = {
    "parallel_corpus": "needs_review",
    "bible_original": "needs_review",
    "bible_iso_improved": "blocked",
    "daily_use_phrases": "needs_review",
}
SUMMARY = {
    "model-a": {"acceptable_95ci": [0.85, 0.95]},
    "model-b": {"acceptable_95ci": [0.40, 0.60]},
}
BIBLE = [
    {"book_id": 1, "chapter": 1, "verse": 1, "reference": "Genesis 1:1",
     "english": "In the beginning.", "zomi_original": "ZOMI_PLACEHOLDER_V1", "zomi_iso": "ZOMI_PLACEHOLDER_ISO_V1"},
    {"book_id": 1, "chapter": 1, "verse": 2, "reference": "Genesis 1:2",
     "english": "And the earth.", "zomi_original": "ZOMI_PLACEHOLDER_V2", "zomi_iso": "ZOMI_PLACEHOLDER_ISO_V2"},
    {"book_id": 1, "chapter": 2, "verse": 1, "reference": "Genesis 2:1",
     "english": "Thus the heavens.", "zomi_original": None, "zomi_iso": "ZOMI_PLACEHOLDER_ISO_V3"},
]
PHRASES = [
    {"id": 1, "english": "Good morning.", "zomi": "ZOMI_PLACEHOLDER_P1", "verified": True},
    {"id": 2, "english": "Thank you.", "zomi": "ZOMI_PLACEHOLDER_P2", "verified": False},
    {"id": 3, "english": "See you tomorrow.", "zomi": "ZOMI_PLACEHOLDER_P3", "verified": True},
]


class ExportTests(unittest.TestCase):
    def setUp(self) -> None:
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.dir = Path(tmp.name)
        corpus = self.dir / "corpus.csv"
        with corpus.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(["index", "en", "zomi", "model", "worker"])
            writer.writerow([0, "The river is wide.", "ZOMI_PLACEHOLDER_CA", "model-a", "w"])
            writer.writerow([1, "My brother cooks rice.", "ZOMI_PLACEHOLDER_CB", "model-b", "w"])
            writer.writerow([2, "I bought 4 eggs.", "ZOMI_PLACEHOLDER_CC", "model-a", "w"])
            writer.writerow([3, "Held out sentence.", "ZOMI_PLACEHOLDER_CD", "model-a", "w"])
        self.database = self.dir / "language.sqlite"
        con = connect(self.database)
        create_schema(con)
        con.execute(
            "INSERT INTO dataset_versions (name, version, created_at, script_version, source_files) "
            "VALUES ('siamsil-language', 'v2', 't', '1.1.0', '{}')"
        )
        with redirect_stdout(StringIO()):
            ingest_parallel(con, corpus, "t")
        held = {"direction": "en-zomi", "source_text": "Held out sentence.", "references": ["x"]}
        apply_corpus_checks(con, [held])
        con.close()
        self.con = sqlite3.connect(self.database)
        self.addCleanup(self.con.close)

    def run_export(self, **overrides) -> tuple[dict, dict[str, list[dict]]]:
        options = dict(
            rights=ALL_NEEDS_REVIEW, allow_uncleared=True, bible=BIBLE, phrases=PHRASES,
            eval_items=[], summary=SUMMARY, min_acceptable=0.8, min_score=0.5,
        )
        options.update(overrides)
        out = self.dir / "out"
        result = export(self.con, out, **options)
        rows = {}
        for split in ("train", "validation", "test"):
            text = (out / f"{split}.jsonl").read_text(encoding="utf-8")
            rows[split] = [json.loads(line) for line in text.splitlines()]
        return result, rows

    def all_rows(self, rows: dict[str, list[dict]]) -> dict[str, dict]:
        return {row["id"]: row for split_rows in rows.values() for row in split_rows}

    def test_nothing_is_exported_until_rights_are_cleared(self) -> None:
        result, rows = self.run_export(allow_uncleared=False)
        self.assertEqual(sum(len(r) for r in rows.values()), 0)
        self.assertEqual(result["uncleared_sources_included"], [])
        self.assertGreater(result["skipped"]["rights:parallel_corpus"], 0)

    def test_blocked_sources_are_never_admitted(self) -> None:
        self.assertNotIn("bible_iso_improved", admitted_sources(ALL_NEEDS_REVIEW, allow_uncleared=True))
        cleared = dict(ALL_NEEDS_REVIEW, daily_use_phrases="cleared")
        self.assertEqual(admitted_sources(cleared, allow_uncleared=False), {"daily_use_phrases"})

    def test_tiers_and_exclusions(self) -> None:
        result, rows = self.run_export()
        by_id = self.all_rows(rows)
        tiers = {row["english"]: row["tier"] for row in by_id.values()}

        self.assertEqual(tiers["The river is wide."], "silver")
        self.assertEqual(tiers["My brother cooks rice."], "bronze")  # model-b below threshold
        self.assertEqual(tiers["I bought 4 eggs."], "bronze")  # number_mismatch flag
        self.assertNotIn("Held out sentence.", tiers)  # eval_holdout split
        self.assertEqual(tiers["In the beginning."], "gold")
        self.assertNotIn("Thus the heavens.", tiers)  # no original ZSF text
        self.assertNotIn("Thank you.", tiers)  # phrase not verified
        self.assertEqual(result["uncleared_sources_included"], ["bible_original", "daily_use_phrases", "parallel_corpus"])
        self.assertEqual(result["silver_models"], ["model-a"])

    def test_bible_uses_original_text_and_keeps_chapters_together(self) -> None:
        _, rows = self.run_export()
        by_id = self.all_rows(rows)
        self.assertEqual(by_id["bible-1-1-1"]["zomi"], "ZOMI_PLACEHOLDER_V1")
        self.assertEqual(by_id["bible-1-1-1"]["split"], by_id["bible-1-1-2"]["split"])
        self.assertNotIn("ZOMI_PLACEHOLDER_ISO_V1", json.dumps(rows))

    def test_evaluation_matches_are_not_exported(self) -> None:
        items = [{"direction": "en-zomi", "source_text": "good MORNING", "references": ["y"]}]
        result, rows = self.run_export(eval_items=items)
        self.assertNotIn("phrase-1", self.all_rows(rows))
        self.assertEqual(result["skipped"]["evaluation_match"], 1)

    def test_split_counts_add_up(self) -> None:
        result, rows = self.run_export()
        for split, split_rows in rows.items():
            self.assertEqual(result["splits"][split]["total"], len(split_rows))
            self.assertEqual(sum(result["splits"][split]["by_tier"].values()), len(split_rows))
            self.assertEqual(sum(result["splits"][split]["by_source"].values()), len(split_rows))

    def test_silver_requires_lower_bound(self) -> None:
        self.assertEqual(silver_models(SUMMARY, 0.8), {"model-a"})
        self.assertEqual(silver_models(None, 0.8), set())
        self.assertEqual(silver_models({"m": {"acceptable_95ci": None}}, 0.0), set())

    def test_main_writes_manifest(self) -> None:
        out = self.dir / "cli"
        with redirect_stdout(StringIO()):
            code = main(["--db", str(self.database), "--out", str(out), "--allow-uncleared"])
        self.assertEqual(code, 0)
        manifest = json.loads((out / "manifest.json").read_text(encoding="utf-8"))
        self.assertEqual(manifest["language_release"], "v2")
        self.assertEqual(manifest["silver_models"], [])
        self.assertIn("parallel_corpus", manifest["uncleared_sources_included"])
        self.assertEqual(set(manifest["files"]), {"train", "validation", "test"})


BIBLE_METADATA = """<DBLMetadata><identification><name>{abbr} name</name><abbreviation>{abbr}</abbreviation>
</identification><language><iso>{iso}</iso><name>x</name></language></DBLMetadata>"""
BIBLE_BOOK = """<usx version="2.0"><book code="{code}" style="id" /><chapter number="1" style="c" />
<para style="p"><verse number="1" style="v" />{text} one<verse number="2-3" style="v" />{text} range
<verse number="4" style="v" />{text} four</para></usx>"""


class BibleVersionExportTests(ExportTests):
    def build_versions(self, versions: dict[str, tuple[str, str]]) -> Path:
        source = self.dir / "usx"
        for abbr, (iso, text) in versions.items():
            usx = source / abbr / "USX_1"
            usx.mkdir(parents=True)
            (source / abbr / "metadata.xml").write_text(BIBLE_METADATA.format(abbr=abbr, iso=iso), encoding="utf-8")
            for code in ("JHN", "MRK"):
                (usx / f"{code}.usx").write_text(BIBLE_BOOK.format(code=code, text=f"{text} {code}"), encoding="utf-8")
        config = self.dir / "bible_versions.json"
        config.write_text(json.dumps({"versions": {"zt": {"name": "Zomi Test", "language_iso": "zom"}}}))
        db = self.dir / "bible_versions.sqlite"
        with redirect_stdout(StringIO()), patch("sys.stderr"):
            import_bibles(["--source", str(source), "--out", str(db), "--manifest", str(self.dir / "m.json"),
                           "--config", str(config), "--rights", str(self.dir / "none.json")])
        return db

    def setUp(self) -> None:
        super().setUp()
        # ZT is configured as Zomi (source code ctd); OTHER is neither Zomi nor English.
        self.versions = self.build_versions({
            "ZT": ("ctd", "ZOMI_PLACEHOLDER"), "EN": ("eng", "English"), "OTHER": ("cnh", "OTHER_PLACEHOLDER"),
        })
        self.rights = dict(ALL_NEEDS_REVIEW, bible_zt="needs_review", bible_en="needs_review", bible_other="needs_review")

    def test_pairs_zomi_with_english_by_book_split(self) -> None:
        result, rows = self.run_export(rights=self.rights, bible_versions=self.versions)
        pairs = {k: v for k, v in self.all_rows(rows).items() if k.startswith("bible-zt-")}
        self.assertEqual(sorted(pairs), ["bible-zt-en-41-1-1", "bible-zt-en-41-1-4", "bible-zt-en-43-1-1", "bible-zt-en-43-1-4"])
        self.assertEqual(pairs["bible-zt-en-43-1-1"]["english"], "English JHN one")
        self.assertEqual(pairs["bible-zt-en-43-1-1"]["split"], "train")
        self.assertEqual(pairs["bible-zt-en-41-1-1"]["split"], "test")  # Mark is a test book
        self.assertEqual(result["skipped"]["bible_verse_range"], 2)
        self.assertNotIn("OTHER_PLACEHOLDER", json.dumps(rows))

    def test_pairs_need_both_versions_cleared(self) -> None:
        rights = dict(self.rights, bible_en="blocked")
        result, rows = self.run_export(rights=rights, bible_versions=self.versions)
        self.assertFalse([k for k in self.all_rows(rows) if k.startswith("bible-zt-")])
        self.assertEqual(result["skipped"]["rights:bible_en"], 4)

    def test_workbook_and_versions_share_book_splits(self) -> None:
        bible = [{"book_english": "Mark", "book_id": 99, "chapter": 1, "verse": 1, "reference": "Mark 1:1",
                  "english": "Workbook Mark.", "zomi_original": "ZOMI_PLACEHOLDER_WB"},
                 {"book_english": "Unknown Book", "book_id": 1, "chapter": 1, "verse": 1, "reference": "x",
                  "english": "Lost.", "zomi_original": "ZOMI_PLACEHOLDER_LOST"}]
        result, rows = self.run_export(rights=self.rights, bible=bible, bible_versions=self.versions)
        by_id = self.all_rows(rows)
        self.assertEqual(by_id["bible-99-1-1"]["split"], by_id["bible-zt-en-41-1-1"]["split"])
        self.assertEqual(result["skipped"]["bible_unknown_book"], 1)


class WorkbookTests(unittest.TestCase):
    def test_loads_bible_and_daily_use_sheets(self) -> None:
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path = Path(tmp.name) / "workbook.xlsx"
        wb = openpyxl.Workbook()
        bible = wb.active
        bible.title = "Bible Data (ISO)"
        bible.append(["Testament", "Book ID", "Book (English)", "Book (Zomi)", "Chapter", "Verse",
                      "KJV English", "Original ZSF Text", "ISO Improved Text", "Changes Applied"])
        bible.append(["OT", 1, "Genesis", "ZOMI_PLACEHOLDER_BOOK", 1, 1,
                      "In the beginning.", "ZOMI_PLACEHOLDER_V1", "ZOMI_PLACEHOLDER_ISO_V1", None])
        daily = wb.create_sheet("Daily Use (Ms. Nem)")
        daily.append(["Daily use phrases"])
        daily.append([])
        daily.append(["#", "Category", "Sub-category", "English", "Zomi", "Notes", "Verified", "Source"])
        daily.append([1, "Greetings", None, "Good morning.", "ZOMI_PLACEHOLDER_P1", None, "Yes", "Teacher"])
        wb.save(path)

        verses, phrases = load_workbook_sources(path)
        self.assertEqual(verses[0]["zomi_original"], "ZOMI_PLACEHOLDER_V1")
        self.assertEqual(verses[0]["book_id"], 1)
        self.assertEqual(phrases[0]["english"], "Good morning.")
        self.assertTrue(phrases[0]["verified"])


if __name__ == "__main__":
    unittest.main()
