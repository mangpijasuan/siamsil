from __future__ import annotations

import csv
import json
import os
import sys
import tempfile
import unittest
from contextlib import contextmanager
from pathlib import Path
from unittest import mock

BACKEND_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_ROOT))
sys.path.insert(0, str(BACKEND_ROOT / "scripts"))
sys.path.insert(0, str(BACKEND_ROOT.parent / "ml_pipeline" / "scripts"))

import dictionary_reviews as dr  # type: ignore  # noqa: E402

# Zomi-side strings are placeholders, never real language data.


def review(sense: str, decision: str, reviewer: str = "R01", correction: str = "", shown: str = "d" * 40) -> dict:
    return {"kind": "review", "sense_id": sense, "decision": decision, "correction": correction, "shown": shown,
            "reviewer": reviewer, "reviewed_on": "2026-11-01"}


def approval(sense: str, decision: str = "approve", approver: str = "E01") -> dict:
    return {"kind": "approval", "sense_id": sense, "decision": decision, "approver": approver,
            "approved_on": "2026-11-02"}


class ConfigTests(unittest.TestCase):
    def test_exactly_one_primary_source(self) -> None:
        import dictionary_review

        config = json.loads((BACKEND_ROOT.parent / "data" / "dictionary_sources.json").read_text(encoding="utf-8"))
        self.assertTrue(all(spec.get("role") in ("primary", "reference") for spec in config["sources"].values()))
        self.assertEqual(dictionary_review.primary_sources(), ["translation_master_50k"])


class RuleTests(unittest.TestCase):
    def test_validate(self) -> None:
        self.assertEqual(dr.validate(review("ZD-000001", "correct")), [])
        self.assertIn("a fix needs the corrected Zomi", dr.validate(review("ZD-000001", "fix")))
        self.assertTrue(dr.validate(review("ZD-000001", "maybe")))
        self.assertTrue(dr.validate(review("X-1", "correct")))
        self.assertTrue(dr.validate(review("ZD-000001", "correct", reviewer="Mangpi")))
        self.assertTrue(dr.validate(approval("ZD-000001", "perhaps")))

    def test_statuses(self) -> None:
        records = [
            review("ZD-000001", "correct"),
            review("ZD-000002", "correct"), approval("ZD-000002"),
            review("ZD-000003", "fix", correction="QA VO"),
            review("ZD-000004", "not_tedim"),
            review("ZD-000005", "correct"), approval("ZD-000005", "reject"),
            review("ZD-000006", "correct"), approval("ZD-000006", approver="R01"),
            approval("ZD-000007"), review("ZD-000007", "correct"),
        ]
        got = {sense: d.status for sense, d in dr.decide(records).items()}
        self.assertEqual(got, {
            "ZD-000001": dr.READY_FOR_APPROVAL,
            "ZD-000002": dr.HUMAN_APPROVED,
            "ZD-000003": dr.NEEDS_CORRECTION,
            "ZD-000004": dr.REJECTED,
            "ZD-000005": dr.REJECTED,
            "ZD-000006": dr.READY_FOR_APPROVAL,  # a reviewer cannot approve their own review
            "ZD-000007": dr.READY_FOR_APPROVAL,  # an approval older than the review does not count
        })
        self.assertEqual(dr.decide(records)["ZD-000003"].corrections, {"R01": "QA VO"})

    def test_latest_review_counts_and_min_reviewers(self) -> None:
        records = [review("ZD-000001", "fix", correction="QA"), review("ZD-000001", "correct")]
        self.assertEqual(dr.decide(records)["ZD-000001"].status, dr.READY_FOR_APPROVAL)
        self.assertEqual(dr.decide(records, min_reviewers=2)["ZD-000001"].status, dr.NEEDS_CORRECTION)

    def test_load_refuses_bad_lines(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "r.jsonl"
            path.write_text(json.dumps(review("ZD-000001", "correct")) + "\n" + "{bad\n", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "line 2"):
                dr.load(path)
            self.assertEqual(dr.load(Path(tmp) / "missing.jsonl"), [])

    def test_revisions(self) -> None:
        revision = {"kind": "revision", "sense_id": "ZD-000001", "before": "QA VO", "after": "QA VU",
                    "proposed_by": ["R01"], "approver": "E01", "approved_on": "2026-11-03"}
        self.assertEqual(dr.validate(revision), [])
        self.assertIn("a lead editor cannot approve their own correction",
                      dr.validate({**revision, "approver": "R01"}))
        self.assertIn("the correction does not change the Zomi", dr.validate({**revision, "after": " QA  VO "}))
        self.assertTrue(dr.validate({**revision, "proposed_by": "R01"}))
        newer = {**revision, "after": "QA VA"}
        self.assertEqual(dr.latest_revisions([revision, newer]), {"ZD-000001": newer})
        # Revisions and declined corrections never count as approvals of a review.
        declined = {"kind": "declined", "sense_id": "ZD-000001", "after": "QA VA", "approver": "E01",
                    "approved_on": "2026-11-03"}
        self.assertEqual(dr.validate(declined), [])
        records = [review("ZD-000001", "correct"), revision, declined]
        self.assertEqual(dr.decide(records)["ZD-000001"].status, dr.READY_FOR_APPROVAL)

    def test_shown_digest_ignores_spacing(self) -> None:
        self.assertEqual(dr.shown_digest("qa  vo"), dr.shown_digest(" qa vo "))
        self.assertNotEqual(dr.shown_digest("qa vo"), dr.shown_digest("qa vu"))


@unittest.skipUnless(os.getenv("SIAMSIL_DATABASE_URL"), "application database is not configured")
class WorkflowTests(unittest.TestCase):
    """Export, record, approve and apply against the database, inside a rolled-back transaction."""

    def setUp(self) -> None:
        from sqlalchemy.orm import Session

        from database import get_database_engine
        from dictionary.repository import DictionaryRepository
        from dictionary_sources import Issue, SourceRead, SourceRecord

        engine = get_database_engine()
        if engine is None:
            self.skipTest("application database is not configured")
        self.connection = engine.connect()
        self.transaction = self.connection.begin()
        self.session = Session(bind=self.connection, join_transaction_mode="create_savepoint")
        records = [
            SourceRecord(f"L{n:06d}", f"zzqreview-{n}", "n.", definition_en=f"Meaning {n}.",
                         equivalents=[("ctd", "gloss", f"QAVO {n}"), ("my", "gloss", f"MY {n}")],
                         issues=[Issue(f"L{n:06d}:v", "Language", "VARIETY_CHECK", "check", None, None, {})],
                         raw={})
            for n in range(1, 5)
        ]
        DictionaryRepository(self.session).import_source(
            "zzq_review_source", title="t", version=None, file_name="t.txt", sha256="b" * 64,
            rights_key="zzq", rights={"display": "needs_review"}, read=SourceRead(records, {}))
        self.session.flush()
        self.tmp = tempfile.TemporaryDirectory()
        self.reviews = Path(self.tmp.name)
        with (self.reviews / "roster.csv").open("w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerows([["code", "role"], ["R01", "reviewer"], ["E01", "lead_editor"]])

        import dictionary_review

        self.tool = dictionary_review

        @contextmanager
        def fake_session():
            yield self.session

        self.patch = mock.patch.object(dictionary_review, "session", fake_session)
        self.patch.start()

    def tearDown(self) -> None:
        self.patch.stop()
        self.tmp.cleanup()
        self.session.close()
        self.transaction.rollback()
        self.connection.close()

    def run_tool(self, *args: str) -> int:
        return self.tool.main(["--reviews-dir", str(self.reviews), *args])

    def fill(self, path: Path, sheet: str, decisions: list[tuple[str, str]]) -> None:
        import openpyxl

        wb = openpyxl.load_workbook(path)
        ws = wb[sheet]
        header = [c.value for c in ws[1]]
        decision_col = header.index("Decision") + 1
        fix_col = header.index("Corrected Zomi") + 1 if "Corrected Zomi" in header else None
        for row, (decision, correction) in zip(range(2, ws.max_row + 1), decisions):
            ws.cell(row, decision_col).value = decision or None
            if fix_col and correction:
                ws.cell(row, fix_col).value = correction
        wb.save(path)

    def status_of(self) -> dict[str, tuple[str, str]]:
        from sqlalchemy import text

        rows = self.session.execute(text(
            "SELECT e.headword, s.review_status, s.status FROM dictionary_senses s "
            "JOIN dictionary_entries e ON e.id = s.entry_id WHERE s.source_id = 'zzq_review_source'")).all()
        return {headword: (review_status, status) for headword, review_status, status in rows}

    def test_full_review_cycle(self) -> None:
        sheet = self.reviews / "batch.xlsx"
        self.assertEqual(self.run_tool("export", "--reviewer", "R01", "--out", str(sheet),
                                       "--source", "zzq_review_source", "--issue", "VARIETY_CHECK"), 0)
        ledger = [json.loads(l) for l in (self.reviews / "dictionary-batches.jsonl").read_text().splitlines()]
        self.assertEqual((ledger[0]["batch"], len(ledger[0]["sense_ids"])), ("dict-R01-0001", 4))

        # A second export does not hand out the same senses.
        again = self.reviews / "again.xlsx"
        self.run_tool("export", "--reviewer", "R01", "--out", str(again), "--source", "zzq_review_source")
        self.assertFalse(again.exists())

        self.fill(sheet, "Review", [("correct", ""), ("fix", "QAVO FIXED"), ("not_tedim", ""), ("", "")])
        with self.assertRaises(SystemExit):  # wrong reviewer for this sheet
            self.run_tool("record", "--sheet", str(sheet), "--reviewer", "R02")
        self.assertEqual(self.run_tool("record", "--sheet", str(sheet), "--reviewer", "R01"), 0)
        self.assertEqual(len(dr.load(self.reviews / "dictionary-reviews.jsonl")), 3)

        approve = self.reviews / "approve.xlsx"
        self.assertEqual(self.run_tool("approval-sheet", "--out", str(approve)), 0)
        self.fill(approve, "Approve", [("approve", "")])
        with self.assertRaises(SystemExit):  # a reviewer cannot approve
            self.run_tool("approve", "--sheet", str(approve), "--approver", "R01")
        self.assertEqual(self.run_tool("approve", "--sheet", str(approve), "--approver", "E01"), 0)

        self.assertEqual(self.run_tool("apply"), 0)
        self.assertEqual(self.status_of(), {
            "zzqreview-1": ("approved", "approved"),
            "zzqreview-2": ("needs_correction", "in_review"),
            "zzqreview-3": ("rejected", "rejected"),
            "zzqreview-4": ("unreviewed", "draft"),
        })

    def zomi_of(self, headword: str) -> str:
        from sqlalchemy import text

        return self.session.execute(text(
            "SELECT string_agg(q.text, ' ; ' ORDER BY q.position) FROM dictionary_equivalents q "
            "JOIN dictionary_senses s ON s.id = q.sense_id JOIN dictionary_entries e ON e.id = s.entry_id "
            "WHERE e.headword = :headword AND q.language = 'ctd'"), {"headword": headword}).scalar()

    def test_a_term_sheet_and_an_approved_correction(self) -> None:
        terms = self.reviews / "terms.json"
        terms.write_text(json.dumps({"terms": {"zzq": {
            "english": ["zzqreview"], "zomi": [{"word": "QAVU", "meaning": "placeholder"}],
            "suggest": [{"from": "QAVO", "to": "QAVU"}], "decided_by": "project owner", "decided_on": "2026-10-31",
        }}}), encoding="utf-8")
        sheet = self.reviews / "term.xlsx"
        self.assertEqual(self.run_tool("export", "--reviewer", "R01", "--out", str(sheet), "--source",
                                       "zzq_review_source", "--term", "zzq", "--terms-file", str(terms)), 0)
        import openpyxl

        ws = openpyxl.load_workbook(sheet)["Review"]
        header = [c.value for c in ws[1]]
        self.assertLess(header.index("Suggested Zomi (AI draft)"), header.index("Decision"))
        self.assertEqual(ws.cell(2, header.index("Suggested Zomi (AI draft)") + 1).value, "QAVU 1")
        with self.assertRaises(SystemExit):
            self.run_tool("export", "--reviewer", "R01", "--out", str(sheet), "--term", "nope",
                          "--terms-file", str(terms))

        self.fill(sheet, "Review", [("fix", "QAVU 1"), ("fix", "QAVU 2"), ("correct", ""), ("", "")])
        self.assertEqual(self.run_tool("record", "--sheet", str(sheet), "--reviewer", "R01"), 0)

        corrections = self.reviews / "corrections.xlsx"
        self.assertEqual(self.run_tool("corrections-sheet", "--out", str(corrections)), 0)
        self.fill(corrections, "Corrections", [("approve", ""), ("decline", "")])
        with self.assertRaises(SystemExit):  # the reviewer who proposed it cannot approve it
            self.run_tool("revise", "--sheet", str(corrections), "--approver", "R01")
        self.assertEqual(self.run_tool("revise", "--sheet", str(corrections), "--approver", "E01"), 0)
        self.assertEqual(self.run_tool("apply"), 0)

        self.assertEqual(self.zomi_of("zzqreview-1"), "QAVU 1")
        self.assertEqual(self.zomi_of("zzqreview-2"), "QAVO 2")  # declined: unchanged
        status = self.status_of()
        self.assertEqual(status["zzqreview-1"], ("approved", "approved"))
        self.assertEqual(status["zzqreview-2"], ("needs_correction", "in_review"))
        # A declined correction is not offered again, and applying twice changes nothing more.
        again = self.reviews / "again.xlsx"
        self.run_tool("corrections-sheet", "--out", str(again))
        self.assertFalse(again.exists())
        self.assertEqual(self.run_tool("apply"), 0)
        self.assertEqual((self.zomi_of("zzqreview-1"), self.status_of()["zzqreview-1"]),
                         ("QAVU 1", ("approved", "approved")))

    def test_reviews_of_changed_zomi_are_ignored(self) -> None:
        from sqlalchemy import text

        sheet = self.reviews / "batch.xlsx"
        self.run_tool("export", "--reviewer", "R01", "--out", str(sheet), "--source", "zzq_review_source")
        self.fill(sheet, "Review", [("correct", "")])
        self.run_tool("record", "--sheet", str(sheet), "--reviewer", "R01")
        self.session.execute(text(
            "UPDATE dictionary_equivalents SET text = 'CHANGED' WHERE language = 'ctd' AND sense_id IN "
            "(SELECT s.id FROM dictionary_senses s JOIN dictionary_entries e ON e.id = s.entry_id "
            "WHERE e.headword = 'zzqreview-1')"))
        self.run_tool("apply")
        self.assertEqual(self.status_of()["zzqreview-1"], ("unreviewed", "draft"))


if __name__ == "__main__":
    unittest.main()
