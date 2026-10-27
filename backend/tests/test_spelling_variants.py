from __future__ import annotations

import csv
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
SCRIPT = ROOT / "ml_pipeline" / "scripts" / "spelling_variants.py"
sys.path.insert(0, str(SCRIPT.parent))

import spelling_variants as sv  # noqa: E402

# Placeholder words only: these tests check counting, not Zomi.
CHECKS = {"checks": [
    {"id": "join", "topic": "joined or split", "zsf": "joined", "zokam": "No explicit rule.", "primer": "split",
     "variants": [["qa qo", "qaqo"]]},
    {"id": "possessive", "topic": "ii or apostrophe", "zsf": "ii", "zokam": "-", "primer": "apostrophe",
     "variants": [["ii", "'"]], "special": "possessive"},
]}


class TokenTests(unittest.TestCase):
    def test_words_keep_inner_marks_and_a_final_apostrophe(self) -> None:
        self.assertEqual(
            sv.tokens("Xa’ yo-zu, “Qe’n” 12 WU' ‘vi"),
            ["xa'", "yo-zu", "qe'n", "wu'", "vi"],
        )

    def test_family_key_ignores_marks_and_doubled_vowels(self) -> None:
        self.assertEqual(sv.family_key("baak"), sv.family_key("bak"))
        self.assertEqual(sv.family_key("ba-ak"), sv.family_key("bak"))
        self.assertEqual(sv.family_key("b'aak"), "bak")
        self.assertNotEqual(sv.family_key("bak"), sv.family_key("bek"))

    def test_phrase_counts(self) -> None:
        source = sv.Source("s", "test")
        source.add("qa qo zu. qaqo zu qa qo zu")
        self.assertEqual((source.count("qa qo"), source.count("qaqo"), source.count("qa qo zu")), (2, 1, 2))
        with self.assertRaises(ValueError):
            source.count("a b c d")


class AnalysisTests(unittest.TestCase):
    def setUp(self) -> None:
        self.one = sv.Source("one", "first")
        self.one.add("qaqo qaqo qa qo ve-ri baak bak bak bak mo' mo mo zi ii")
        self.two = sv.Source("two", "second")
        self.two.add("qa qo qa qo veri veri mo' mo' bak zi' zi'")

    def test_joins_find_all_three_spellings(self) -> None:
        rows = {row["joined"]: row for row in sv.joins([self.one, self.two], min_total=1)}
        self.assertEqual(rows["qaqo"]["split"], "qa qo")
        self.assertEqual(rows["qaqo"]["total"], 2 + 3)
        self.assertEqual(rows["veri"]["total"], 2 + 1)  # joined twice, hyphenated once
        self.assertEqual(rows["qaqo"]["minority"], 2)

    def test_families_fold_the_possessive_apostrophe(self) -> None:
        families = {tuple(f["forms"]): f for f in sv.variant_families([self.one, self.two], min_total=1)}
        self.assertIn(("bak", "baak"), families)
        # "mo" and "mo'" are one word (the apostrophe is grammar), so they are not a family.
        self.assertFalse(any("mo" in forms for forms in families))

    def test_rule_checks_count_each_variant_per_source(self) -> None:
        rows = sv.rule_checks([self.one, self.two], CHECKS["checks"])
        counts = {(row["check"], row["variant"]): row["counts"] for row in rows}
        self.assertEqual(counts[("join", "qa qo")], {"one": 1, "two": 2})
        self.assertEqual(counts[("join", "qaqo")], {"one": 2, "two": 0})
        self.assertEqual(counts[("possessive", "ii")], {"one": 1, "two": 0})
        self.assertEqual(counts[("possessive", "word ending in an apostrophe")], {"one": 1, "two": 4})


class SuggestionTests(unittest.TestCase):
    def test_the_spelling_most_sources_prefer_is_suggested(self) -> None:
        counts = {"one": {"qaqo": 9, "qa qo": 1}, "two": {"qaqo": 5}, "three": {"qa qo": 4}}
        self.assertEqual(sv.suggest(counts, ["qaqo", "qa qo"], "three"),
                         ("qaqo", "mixed", "preferred by one, two; other spellings in three (qa qo)"))
        self.assertEqual(sv.suggest({"one": {"qaqo": 9}, "two": {"qaqo": 3}}, ["qaqo", "qa qo"], "one")[:2],
                         ("qaqo", "strong"))

    def test_the_reference_source_breaks_a_tie(self) -> None:
        counts = {"one": {"bak": 5}, "two": {"baak": 5}}
        self.assertEqual(sv.suggest(counts, ["bak", "baak"], "two")[0], "baak")
        self.assertEqual(sv.suggest(counts, ["bak", "baak"], "one")[0], "bak")

    def test_no_suggestion_when_a_source_uses_both_often_or_too_few_uses(self) -> None:
        best, strength, why = sv.suggest({"one": {"tu": 30, "tuu": 10}}, ["tu", "tuu"], "one")
        self.assertEqual((best, strength), ("", "check meaning"))
        self.assertIn("tu (30) and tuu (10)", why)
        self.assertEqual(sv.suggest({"one": {"tu": 2}}, ["tu", "tuu"], "one")[:2], ("", "too few uses"))


class CommandTests(unittest.TestCase):
    def test_reports_from_text_and_workbook_sources(self) -> None:
        from openpyxl import Workbook, load_workbook

        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            (folder / "a.txt").write_text("qaqo qa qo bak baak bak\nqa qo zi ii\n", encoding="utf-8")
            book = Workbook()
            book.active.append(["ID", "Text"])
            book.active.append([1, "qa qo bak zi'"])
            book.active.append([2, None])
            book.save(folder / "d.xlsx")
            checks = folder / "checks.json"
            checks.write_text(json.dumps(CHECKS), encoding="utf-8")
            out = folder / "out"
            result = subprocess.run(
                [sys.executable, str(SCRIPT), "--text", f"book={folder / 'a.txt'}",
                 "--xlsx", f"dict={folder / 'd.xlsx'}:Text", "--checks", str(checks), "--min-total", "1", "--out", str(out)],
                capture_output=True, text=True,
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            with (out / "rule_checks.csv").open(encoding="utf-8") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual([(r["variant"], r["book"], r["dict"]) for r in rows if r["check"] == "join"],
                             [("qa qo", "2", "1"), ("qaqo", "1", "0")])
            with (out / "forms.csv").open(encoding="utf-8") as handle:
                forms = {row["form"]: row for row in csv.DictReader(handle)}
            self.assertEqual((forms["bak"]["book"], forms["bak"]["dict"], forms["bak"]["total"]), ("2", "1", "3"))
            summary = json.loads((out / "summary.json").read_text(encoding="utf-8"))
            self.assertEqual([s["name"] for s in summary["sources"]], ["book", "dict"])
            workbook = load_workbook(out / "Siamsil_Zolai_Standard_Review.xlsx")
            self.assertEqual(workbook.sheetnames, ["Read me", "Rules", "Joins", "Families", "Sources"])
            rules = workbook["Rules"]
            self.assertEqual(rules["A2"].value, "join")
            header = [cell.value for cell in rules[1]]
            self.assertIn("Decision", header)
            self.assertLess(header.index("Siamsil suggestion (AI)"), header.index("Decision"))

    def test_unknown_column_and_missing_sources_are_refused(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            book = Workbook()
            book.active.append(["ID", "Text"])
            book.save(folder / "d.xlsx")
            (folder / "a.txt").write_text("qa qo\n", encoding="utf-8")
            for args in ([], ["--xlsx", f"dict={folder / 'd.xlsx'}:Nope"],
                         ["--text", f"book={folder / 'a.txt'}", "--reference", "nobody"]):
                with self.subTest(args=args):
                    result = subprocess.run([sys.executable, str(SCRIPT), *args, "--out", str(folder / "out")],
                                            capture_output=True, text=True)
                    self.assertNotEqual(result.returncode, 0)


if __name__ == "__main__":
    unittest.main()
