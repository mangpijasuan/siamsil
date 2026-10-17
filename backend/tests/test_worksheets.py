from __future__ import annotations

import csv
import json
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from io import StringIO
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT_ROOT / "ml_pipeline" / "scripts"))

import openpyxl  # noqa: E402

from eval_set import load_items, validate_items  # type: ignore  # noqa: E402
from worksheets import main  # type: ignore  # noqa: E402

# Zomi-side strings are placeholders, never real language data.


def workbook(path: Path, sheet: str, header: list[str], rows: list[list]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Instructions"
    data = wb.create_sheet(sheet)
    data.append(header)
    for row in rows:
        data.append(row)
    wb.save(path)


def run(*argv: str) -> str:
    out = StringIO()
    with redirect_stdout(out):
        main(list(argv))
    return out.getvalue()


class TestSetRoundTripTests(unittest.TestCase):
    def setUp(self) -> None:
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.dir = Path(tmp.name)
        self.sheet = self.dir / "translate.xlsx"
        workbook(self.sheet, "Translate", ["ID", "Topic", "English", "Zomi translation", "Notes"], [
            ["en-zomi-0001", "health", "My stomach hurts after eating the fish.", "ZOMI_PLACEHOLDER_1", ""],
            ["en-zomi-0002", "news", "No one in the family was injured.", "ZOMI_PLACEHOLDER_2", "two options"],
            ["en-zomi-0003", "informal", "Wait for me at the bus stop!", None, ""],
        ])

    def test_import_review_and_merge(self) -> None:
        draft = self.dir / "draft.jsonl"
        self.assertIn("2 translated, 1 not yet translated",
                      run("testset-import", str(self.sheet), "--translator", "T01", "--origin", "written for Siamsil", "--out", str(draft)))

        review = self.dir / "review.xlsx"
        run("review-sheet", str(draft), "--out", str(review))
        wb = openpyxl.load_workbook(review)
        ws = wb["Review"]
        self.assertEqual([c.value for c in ws[1]][:5], ["ID", "Topic", "English", "Zomi translation", "Decision"])
        ws["E2"] = "Accept"
        ws["E3"] = "Correct"
        ws["F3"] = "ZOMI_PLACEHOLDER_2_FIXED"
        wb.save(review)

        final = self.dir / "en-zomi-v1.jsonl"
        self.assertIn("2 items (1 accepted, 1 corrected)",
                      run("review-merge", str(draft), str(review), "--reviewer", "R01", "--out", str(final)))
        items, parse_errors = load_items(final)
        errors, _ = validate_items(items)
        self.assertEqual((parse_errors, errors), ([], []))
        self.assertEqual(items[1]["references"], ["ZOMI_PLACEHOLDER_2_FIXED"])
        self.assertEqual({item["reviewer"] for item in items}, {"R01"})

    def test_reviewer_must_differ_from_translator(self) -> None:
        draft = self.dir / "draft.jsonl"
        run("testset-import", str(self.sheet), "--translator", "T01", "--origin", "x", "--out", str(draft))
        review = self.dir / "review.xlsx"
        run("review-sheet", str(draft), "--out", str(review))
        wb = openpyxl.load_workbook(review)
        wb["Review"]["E2"] = "Accept"
        wb.save(review)
        with self.assertRaises(SystemExit), redirect_stdout(StringIO()):
            main(["review-merge", str(draft), str(review), "--reviewer", "T01", "--out", str(self.dir / "x.jsonl")])
        self.assertFalse((self.dir / "x.jsonl").exists())

    def test_correct_without_text_is_rejected(self) -> None:
        draft = self.dir / "draft.jsonl"
        run("testset-import", str(self.sheet), "--translator", "T01", "--origin", "x", "--out", str(draft))
        review = self.dir / "review.xlsx"
        run("review-sheet", str(draft), "--out", str(review))
        wb = openpyxl.load_workbook(review)
        wb["Review"]["E2"] = "Correct"
        wb.save(review)
        with self.assertRaises(SystemExit), redirect_stdout(StringIO()):
            main(["review-merge", str(draft), str(review), "--reviewer", "R01", "--out", str(self.dir / "x.jsonl")])


class RatingsImportTests(unittest.TestCase):
    def test_rate_sheet_becomes_review_csv(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            sheet = Path(tmp) / "rate.xlsx"
            workbook(sheet, "Rate", ["ID", "English", "Zomi (machine-made)", "Meaning (0-2)", "Fluency (0-2)", "Notes"], [
                ["R00001", "Is the shop open?", "ZOMI_PLACEHOLDER_A", 2, 1.0, ""],
                ["R00002", "It is late.", "ZOMI_PLACEHOLDER_B", None, None, "unsure"],
            ])
            out = Path(tmp) / "review.csv"
            self.assertIn("1 of 2 rows rated", run("ratings-import", str(sheet), "--out", str(out)))
            rows = list(csv.DictReader(out.open(encoding="utf-8")))
            self.assertEqual((rows[0]["sample_id"], rows[0]["meaning"], rows[0]["fluency"]), ("R00001", "2", "1"))
            self.assertEqual(rows[1]["meaning"], "")


if __name__ == "__main__":
    unittest.main()
