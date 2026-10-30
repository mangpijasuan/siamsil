from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT_ROOT / "ml_pipeline" / "scripts"))

import openpyxl  # noqa: E402

from dictionary_sources import (  # type: ignore  # noqa: E402
    READERS,
    normalize_pos,
    parse_ordbog_pages,
    read_learners_json,
    read_master_workbook,
    read_standard_workbook,
    read_tongdot_tsv,
    split_alternatives,
)

# Zomi-side strings are placeholders, never real language data.

MASTER_HEADER = ["Entry ID", "Headword", "Part of speech", "Sense number", "Simple English definition",
                 "Current Standard Zomi", "Current Burmese (reference only)", "Usage tier code", "Domain",
                 "Overall decision", "Correction required?"]
ISSUE_HEADER = ["Item ID", "Record ID", "Issue category", "Issue code", "Description", "Priority",
                "Required reviewer", "Record link"]


class ReaderTests(unittest.TestCase):
    def setUp(self) -> None:
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.dir = Path(tmp.name)

    def test_master_workbook_keeps_ids_splits_alternatives_and_links_issues(self) -> None:
        wb = openpyxl.Workbook()
        records = wb.active
        records.title = "01_All_Records"
        records.append(["IDENTIFICATION (group header row)"])
        records.append(MASTER_HEADER)
        records.append(["ZD-000007", "abandon", "v.", "1", "To leave.", "ZOMI_A1; ZOMI_A2", "MY_A", "A",
                        "unclassified", "Not reviewed", "=chk_correction_required"])
        records.append(["ZD-000009", "Abandon", "n.", "1", "Freedom.", "ZOMI_B", None, "B", "unclassified",
                        "Not reviewed", "=chk"])
        issues = wb.create_sheet("04_Open_Issue_Index")
        issues.append(["OPEN ISSUE INDEX (title row)"])
        issues.append(ISSUE_HEADER)
        issues.append(["OI-000001", "ZD-000007", "Standard Zomi", "Z_X", "check it", "Medium", "Native", "Active record"])
        issues.append(["OI-000002", "", "Recovery", "R_X", "not in master", "Low", "Editor",
                       "Recovery candidate: not in the master"])
        path = self.dir / "master.xlsx"
        wb.save(path)

        read = read_master_workbook(path)
        first = next(r for r in read.records if r.record_id == "ZD-000007")
        self.assertEqual(first.sense_id, "ZD-000007")
        self.assertEqual(first.part_of_speech, "v.")
        self.assertEqual(first.equivalents, [("ctd", "equivalent", "ZOMI_A1"), ("ctd", "equivalent", "ZOMI_A2"),
                                             ("my", "equivalent", "MY_A")])
        self.assertEqual([i.source_issue_id for i in first.issues], ["OI-000001"])
        self.assertEqual([i.source_issue_id for i in read.unlinked_issues], ["OI-000002"])
        self.assertNotIn("Overall decision", first.raw)
        self.assertNotIn("Correction required?", first.raw)

    def test_learners_json_reads_senses_examples_and_phrases(self) -> None:
        path = self.dir / "learners.json"
        path.write_text(json.dumps({"entries": [{
            "id": 12, "headword": "book", "ipa": "x", "part_of_speech": "noun",
            "senses": [{"number": 1, "definition_en": "Pages.", "zomi_equivalents": "ZOMI_C1; ZOMI_C2",
                        "explanation_zomi": "ZOMI_EXPLAIN", "common_phrases": [{"english": "a book", "zomi": "ZOMI_P"}],
                        "examples": [{"english": "I read.", "zomi": "ZOMI_E"}, {"english": "", "zomi": ""}]}],
        }]}), encoding="utf-8")
        record = read_learners_json(path).records[0]
        self.assertEqual(record.record_id, "000012:1")
        self.assertEqual(record.part_of_speech, "n.")
        self.assertEqual(record.explanation_ctd, "ZOMI_EXPLAIN")
        self.assertEqual(len(record.equivalents), 2)
        self.assertEqual(record.examples, [("phrase", "a book", "ZOMI_P"), ("example", "I read.", "ZOMI_E")])

    def test_standard_workbook_turns_flags_into_issues(self) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Master Data"
        ws.append(["ID", "Letter", "Headword", "POS", "POS (full)", "Definition (Zomi)", "Flagged", "Suggested Correction"])
        ws.append([1, "A", "able", "adj.", "adjective", "ZOMI_D", "No", None])
        ws.append([2, "A", "abx", "? n.", None, "ZOMI_E", "Yes", "abc"])
        ws.append([3, "A", None, "n.", "noun", "ZOMI_F", "No", None])
        path = self.dir / "standard.xlsx"
        wb.save(path)
        read = read_standard_workbook(path)
        self.assertEqual([r.record_id for r in read.records], ["1", "2"])
        self.assertEqual(read.records[0].part_of_speech, "adj.")
        self.assertIsNone(read.records[1].part_of_speech, "a questioned label is left unset, not guessed")
        self.assertEqual(read.records[1].issues[0].description, "abc")
        self.assertEqual(read.skipped, {"no_id_or_headword": 1})


class OrdbogTests(unittest.TestCase):
    def test_entries_continuations_and_layouts(self) -> None:
        pages = [
            "Zomi Ordbog\ndansk-[engelsk]-zomi\n",
            "1\nafsky [abhor] (v.): ZOMI_1, ZOMI_2\nZOMI_3\nfarvel [farewell]: ZOMI_4\n"
            "lang dansk etiket, anden [proclaim]\n(v.): ZOMI_5\nførste del af etiketten\n"
            "[contraception] (n.): ZOMI_6\nstart [split]",
            "(n.): ZOMI_7\nintet [---] (n.): ZOMI_8\n",
        ]
        read = parse_ordbog_pages(pages)
        by = {r.headword: r for r in read.records}
        self.assertEqual(sorted(by), ["abhor", "contraception", "farewell", "proclaim", "split"])
        self.assertEqual(by["abhor"].equivalents, [("ctd", "gloss", "ZOMI_1, ZOMI_2 ZOMI_3")])
        self.assertIsNone(by["farewell"].part_of_speech)
        self.assertEqual(by["proclaim"].raw["occurrences"][0]["danish"], "lang dansk etiket, anden")
        self.assertEqual(by["proclaim"].equivalents[0][2], "ZOMI_5")
        self.assertEqual(by["contraception"].raw["occurrences"][0]["danish"], "første del af etiketten")
        self.assertEqual(by["split"].equivalents[0][2], "ZOMI_7")
        self.assertEqual(by["split"].raw["occurrences"][0]["page"], 2)
        self.assertEqual(read.skipped, {"no_english_headword": 1})

    def test_repeats_under_other_danish_labels_become_one_sense(self) -> None:
        pages = ["a [word] (n.): ZOMI_X\nb [word] (n.): ZOMI_X\nc [word] (n.): ZOMI_Y\n", "d [word] (n.): ZOMI_X\n"]
        read = parse_ordbog_pages(pages)
        self.assertEqual(len(read.records), 2)
        first = read.records[0]
        self.assertEqual(first.record_id, "p001-001")
        self.assertEqual([(o["danish"], o["page"]) for o in first.raw["occurrences"]], [("a", 1), ("b", 1), ("d", 2)])
        self.assertEqual(read.skipped["repeated_under_another_danish_label"], 2)
        self.assertEqual(len({r.record_id for r in read.records}), 2)

    def test_joined_entries_are_flagged_not_guessed(self) -> None:
        read = parse_ordbog_pages(["one [first] (v.): ZOMI_Atwo [second] (n.): ZOMI_B\n"])
        self.assertEqual(len(read.records), 1)
        self.assertEqual(read.records[0].issues[0].code, "MERGED_ENTRY")


class HelperTests(unittest.TestCase):
    def test_normalize_pos(self) -> None:
        cases = {"n.": "n.", "noun": "n.", "transitive verb": "v.", "article; determiner": "det.",
                 "exclamation": "interj.", "prop. n.": "prop. n.", "modal verb": "modal v.", "? n.": None,
                 "": None, None: None}
        for value, expected in cases.items():
            self.assertEqual(normalize_pos(value), expected, value)

    def test_split_alternatives(self) -> None:
        self.assertEqual(split_alternatives("A; B ;; C"), ["A", "B", "C"])
        self.assertEqual(split_alternatives(None), [])

    def test_config_sources_have_rights_and_known_formats(self) -> None:
        config = json.loads((PROJECT_ROOT / "data" / "dictionary_sources.json").read_text(encoding="utf-8"))
        rights = json.loads((PROJECT_ROOT / "data" / "rights.json").read_text(encoding="utf-8"))["sources"]
        for source_id, spec in config["sources"].items():
            self.assertIn(spec["format"], READERS, source_id)
            self.assertIn(spec["rights"], rights, source_id)
            self.assertIn(spec["root"], ("data_repo", "project"), source_id)
            self.assertRegex(spec["sha256"], r"^[0-9a-f]{64}$", source_id)


class TongDotTests(unittest.TestCase):
    def test_tongdot_lines_become_labelled_senses(self) -> None:
        lines = [
            "Zomi-English Dictionary (from TongDot.com)", "Source: test", "Entries: 4", "",
            "Qavo\tQavo\t(n.): zu ve, ri ko",
            "Qavo\tqavo\t(n.) \u1000\u1001\u1002\u104b",
            "Zeta\tZeta\t(vt) \u1000\u1001 zeta oneself",
            "Wix\tWix\t((abbrev).) wo ka",
            "Broken line without tabs",
            "Empty\tEmpty\t(n.)",
        ]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "tongdot.txt"
            path.write_text("\n".join(lines), encoding="utf-8")
            read = read_tongdot_tsv(path)
        self.assertEqual([r.record_id for r in read.records], ["L000005", "L000006", "L000007", "L000008"])
        chin, burmese, mixed, abbrev = read.records
        self.assertEqual((chin.headword, chin.part_of_speech, chin.equivalents), ("Qavo", "n.", [("ctd", "gloss", "zu ve, ri ko")]))
        self.assertEqual([i.code for i in chin.issues], ["VARIETY_CHECK"])
        self.assertEqual(burmese.equivalents[0][0], "my")
        self.assertEqual(burmese.issues, [])
        self.assertEqual((mixed.part_of_speech, [i.code for i in mixed.issues]), ("v.", ["MIXED_SCRIPT"]))
        self.assertIsNone(abbrev.part_of_speech)
        self.assertEqual(abbrev.equivalents, [("ctd", "gloss", "wo ka")])
        self.assertEqual(chin.raw["definition"], "(n.): zu ve, ri ko")
        self.assertEqual(read.skipped, {"not_three_columns": 4, "no_definition_text": 1})


if __name__ == "__main__":
    unittest.main()
