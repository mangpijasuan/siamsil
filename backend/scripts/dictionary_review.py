"""Review the editorial dictionary in batches: spreadsheets out, decisions back (docs/DICTIONARY.md).

From the repository root, with SIAMSIL_DATABASE_URL set and the dictionary imported:

  export          a sheet of senses for one reviewer (default 500), remembered so no one else gets them;
                  by default only the primary source's senses (data/dictionary_sources.json), with the
                  reference sources shown beside them
  record          save a reviewer's completed sheet to the decisions file
  approval-sheet  senses every reviewer found correct, for a lead editor
  approve         save the lead editor's decisions
  apply           set each sense's review status in the database from the decisions file
  status          counts per status

  python backend/scripts/dictionary_review.py export --reviewer R01 --out batch-R01-1.xlsx
  python backend/scripts/dictionary_review.py record --sheet batch-R01-1.xlsx --reviewer R01
  python backend/scripts/dictionary_review.py approval-sheet --out approve-1.xlsx
  python backend/scripts/dictionary_review.py approve --sheet approve-1.xlsx --approver E01
  python backend/scripts/dictionary_review.py apply

Decisions, the batch ledger, and the roster live in the private siamsil-data repository
(reviews/dictionary-reviews.jsonl, reviews/dictionary-batches.jsonl, reviews/roster.csv); names
never appear in them, only roster codes. Sheets hold dictionary text: keep them private.
"""

from __future__ import annotations

import argparse
import json
import sys
from datetime import date, datetime, timezone
from pathlib import Path

BACKEND_ROOT = Path(__file__).resolve().parents[1]
PROJECT_ROOT = BACKEND_ROOT.parent
sys.path.insert(0, str(BACKEND_ROOT))
sys.path.insert(0, str(PROJECT_ROOT / "ml_pipeline" / "scripts"))

import dictionary_reviews as dr  # type: ignore  # noqa: E402
from data_repo import default_data_repo  # type: ignore  # noqa: E402

TIER_ORDER = "CASE usage_tier WHEN 'A1' THEN 1 WHEN 'A2' THEN 2 WHEN 'B1' THEN 3 WHEN 'B2' THEN 4 " \
             "WHEN 'C1' THEN 5 WHEN 'C2' THEN 6 ELSE 7 END"
REVIEW_COLUMNS = ["Sense ID", "Headword", "Part of speech", "English meaning", "Zomi to check",
                  "Burmese (reference)", "Example", "Other sources say", "Source", "Flags",
                  "Decision", "Corrected Zomi", "Notes", "Shown"]
APPROVE_COLUMNS = ["Sense ID", "Headword", "Part of speech", "English meaning", "Zomi", "Reviewed by",
                   "Decision", "Notes"]
INSTRUCTIONS = (
    "Siamsil Dictionary review",
    "",
    "For every row, read the English meaning and the Zomi to check. Choose a Decision:",
    "• correct: the Zomi is right and natural Tedim Zolai for this meaning.",
    "• fix: something should change. Write the full corrected Zomi in 'Corrected Zomi'.",
    "• not_tedim: it is another Chin language (Hakha, Falam, Mizo, ...) or not Zomi at all.",
    "• wrong: the Zomi does not mean this.",
    "Leave Decision blank if you are unsure; the row stays open for someone else.",
    "'Other sources say' shows what other dictionaries give for the same word. It may help, but it may also be wrong.",
    "'Flags' are problems the sources reported (VARIETY_CHECK: the source mixes Chin languages).",
    "Do not use Google Translate, ChatGPT, Gemini, or other machine translation.",
    "Do not change the Sense ID or Shown columns. Keep this file private.",
)


def now_utc() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def primary_sources(config: Path = PROJECT_ROOT / "data" / "dictionary_sources.json") -> list[str]:
    sources = json.loads(config.read_text(encoding="utf-8"))["sources"]
    return [source_id for source_id, spec in sources.items() if spec.get("role") == "primary"]


def default_reviews_dir() -> Path:
    return default_data_repo() / "reviews"


def session():
    from database import get_session_factory

    return get_session_factory()()


# Reading the database ---------------------------------------------------------------------------

SENSE_SQL = """
SELECT s.id, e.id AS entry_id, e.headword, e.part_of_speech, s.definition_en, s.explanation_ctd,
       s.source_id, s.review_status,
       (SELECT string_agg(q.text, ' ; ' ORDER BY q.position) FROM dictionary_equivalents q
         WHERE q.sense_id = s.id AND q.language = 'ctd') AS ctd,
       (SELECT string_agg(q.text, ' ; ' ORDER BY q.position) FROM dictionary_equivalents q
         WHERE q.sense_id = s.id AND q.language = 'my') AS my,
       (SELECT coalesce(x.english, '') || ' → ' || coalesce(x.ctd, '') FROM dictionary_examples x
         WHERE x.sense_id = s.id ORDER BY x.position LIMIT 1) AS example,
       (SELECT string_agg(DISTINCT i.code, ', ') FROM dictionary_issues i
         WHERE i.sense_id = s.id AND i.status = 'open') AS flags,
       (SELECT count(DISTINCT o.source_id) FROM dictionary_senses o WHERE o.entry_id = e.id) AS sources,
       s.usage_tier
FROM dictionary_senses s JOIN dictionary_entries e ON e.id = s.entry_id
"""


def zomi_of(row) -> str:
    """The Zomi a reviewer checks: the equivalents, then any Zomi explanation."""
    parts = [part for part in (row.ctd, row.explanation_ctd) if part]
    return " | ".join(parts)


def current_digests(db, sense_ids: list[str]) -> dict[str, str]:
    from sqlalchemy import text

    if not sense_ids:
        return {}
    rows = db.execute(text(SENSE_SQL + " WHERE s.id = ANY(:ids)"), {"ids": sense_ids}).all()
    return {row.id: dr.shown_digest(zomi_of(row)) for row in rows}


def other_sources(db, entry_id: str, sense_id: str) -> str:
    from sqlalchemy import text

    rows = db.execute(text(
        "SELECT s.source_id, string_agg(q.text, ' ; ' ORDER BY q.position) AS ctd "
        "FROM dictionary_senses s JOIN dictionary_equivalents q ON q.sense_id = s.id AND q.language = 'ctd' "
        "WHERE s.entry_id = :entry AND s.id <> :sense GROUP BY s.id, s.source_id ORDER BY s.id LIMIT 4"),
        {"entry": entry_id, "sense": sense_id}).all()
    return "\n".join(f"{row.source_id}: {row.ctd}" for row in rows)


# Ledger of exported batches ---------------------------------------------------------------------

def load_batches(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with path.open(encoding="utf-8") as handle:
        return [json.loads(line) for line in handle if line.strip()]


def next_batch_name(batches: list[dict], prefix: str) -> str:
    taken = {batch["batch"] for batch in batches}
    number = 1
    while f"{prefix}-{number:04d}" in taken:
        number += 1
    return f"{prefix}-{number:04d}"


# Spreadsheets -------------------------------------------------------------------------------------

def write_sheet(path: Path, title: str, columns: list[str], rows: list[list], decisions: tuple[str, ...],
                about: dict, decision_column: str, editable: tuple[str, ...], hidden: tuple[str, ...] = ()) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    intro = wb.active
    intro.title = "Instructions"
    intro.column_dimensions["A"].width = 115
    for line in INSTRUCTIONS:
        intro.append([line])
    intro["A1"].font = Font(bold=True, size=14, color="1B2A4A")

    sheet = wb.create_sheet(title)
    sheet.append(columns)
    for index, name in enumerate(columns, start=1):
        cell = sheet.cell(1, index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1B2A4A")
        width = 12 if name in ("Sense ID", "Part of speech", "Decision", "Source") else 34
        sheet.column_dimensions[cell.column_letter].width = width
        if name in hidden:
            sheet.column_dimensions[cell.column_letter].hidden = True
    sheet.freeze_panes = "C2"
    editable_index = [columns.index(name) + 1 for name in editable]
    for row in rows:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"  # dictionary text is never a formula
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for index in editable_index:
            sheet.cell(sheet.max_row, index).fill = PatternFill("solid", fgColor="FFF4C2")
    choice = DataValidation(type="list", formula1='"' + ",".join(decisions) + '"', allow_blank=True)
    sheet.add_data_validation(choice)
    if rows:
        letter = sheet.cell(1, columns.index(decision_column) + 1).column_letter
        choice.add(f"{letter}2:{letter}{sheet.max_row}")

    info = wb.create_sheet("About")
    for key, value in about.items():
        info.append([key, value])
    info.sheet_state = "hidden"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def read_sheet(path: Path, title: str) -> tuple[list[dict], dict]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if title not in wb.sheetnames or "About" not in wb.sheetnames:
        sys.exit(f"{path} is not a Siamsil dictionary {title.lower()} sheet")
    about = {str(k): ("" if v is None else str(v)) for k, v, *_ in wb["About"].iter_rows(values_only=True)}
    raw = list(wb[title].iter_rows(values_only=True))
    wb.close()
    header = [str(h or "").strip() for h in raw[0]] if raw else []
    rows = []
    for values in raw[1:]:
        if not any(v not in (None, "") for v in values):
            continue
        rows.append({name: ("" if v is None else str(v).strip()) for name, v in zip(header, values)})
    return rows, about


# Commands -----------------------------------------------------------------------------------------

def cmd_export(args: argparse.Namespace) -> int:
    from sqlalchemy import text

    reviews_dir = args.reviews_dir
    roster = dr.load_roster(reviews_dir / "roster.csv") if (reviews_dir / "roster.csv").exists() else None
    if roster is not None and roster.get(args.reviewer) != "reviewer":
        sys.exit(f"{args.reviewer} is not listed as a reviewer in {reviews_dir / 'roster.csv'}")
    batches = load_batches(reviews_dir / "dictionary-batches.jsonl")
    reviewed = {record["sense_id"] for record in dr.load(reviews_dir / "dictionary-reviews.jsonl")}
    taken = set() if args.again else {sense for batch in batches for sense in batch["sense_ids"]}

    where = ["s.review_status = 'unreviewed'", "s.status NOT IN ('rejected', 'deferred')",
             "(EXISTS (SELECT 1 FROM dictionary_equivalents q WHERE q.sense_id = s.id AND q.language = 'ctd') "
             "OR s.explanation_ctd IS NOT NULL)"]
    params: dict = {}
    if args.source:
        where.append("s.source_id = :source")
        params["source"] = args.source
    elif not args.all_sources:
        where.append("s.source_id = ANY(:primary)")
        params["primary"] = primary_sources()
    if args.issue:
        where.append("EXISTS (SELECT 1 FROM dictionary_issues i WHERE i.sense_id = s.id AND i.status = 'open' "
                     "AND i.code = :issue)")
        params["issue"] = args.issue
    order = f"sources DESC, {TIER_ORDER.replace('usage_tier', 's.usage_tier')}, e.headword, s.id"
    with session() as db:
        candidates = db.execute(text(f"{SENSE_SQL} WHERE {' AND '.join(where)} ORDER BY {order}"), params)
        chosen = []
        for row in candidates:
            if row.id in taken or row.id in reviewed:
                continue
            chosen.append(row)
            if len(chosen) >= args.size:
                break
        rows = []
        for row in chosen:
            zomi = zomi_of(row)
            rows.append([row.id, row.headword, row.part_of_speech or "", row.definition_en or "", zomi,
                         row.my or "", row.example or "", other_sources(db, row.entry_id, row.id),
                         row.source_id, row.flags or "", None, "", "", dr.shown_digest(zomi)])
    if not rows:
        print("Nothing left to export with these filters.")
        return 0
    name = args.batch or next_batch_name(batches, f"dict-{args.reviewer}")
    about = {"kind": "dictionary-review", "batch": name, "reviewer": args.reviewer, "exported_at": now_utc()}
    write_sheet(args.out, "Review", REVIEW_COLUMNS, rows, dr.REVIEW_DECISIONS, about,
                "Decision", ("Decision", "Corrected Zomi", "Notes"), hidden=("Shown",))
    dr.append(reviews_dir / "dictionary-batches.jsonl", [{
        "kind": "batch", "batch": name, "reviewer": args.reviewer, "exported_at": about["exported_at"],
        "sense_ids": [row[0] for row in rows]}])
    print(f"Wrote {len(rows)} senses to {args.out} (batch {name}, reviewer {args.reviewer})")
    return 0


def cmd_record(args: argparse.Namespace) -> int:
    rows, about = read_sheet(args.sheet, "Review")
    if about.get("kind") != "dictionary-review":
        sys.exit(f"{args.sheet} is not a dictionary review sheet")
    if about.get("reviewer") != args.reviewer:
        sys.exit(f"{args.sheet} was exported for {about.get('reviewer')}, not {args.reviewer}")
    path = args.reviews_dir / "dictionary-reviews.jsonl"
    roster = dr.load_roster(args.reviews_dir / "roster.csv")
    records, problems = [], []
    for number, row in enumerate(rows, start=2):
        decision = row.get("Decision", "").strip().lower()
        if not decision:
            continue
        record = {"kind": "review", "sense_id": row.get("Sense ID", ""), "decision": decision,
                  "correction": row.get("Corrected Zomi", ""), "shown": row.get("Shown", ""),
                  "reviewer": args.reviewer, "reviewed_on": args.date, "recorded_at": now_utc(),
                  "batch": about.get("batch", ""), "notes": row.get("Notes", "")}
        errors = dr.validate(record)
        if errors:
            problems.append(f"row {number}: {'; '.join(errors)}")
        records.append(record)
    problems += dr.roster_problems(records, roster)
    if problems:
        sys.exit("Nothing was saved:\n  " + "\n  ".join(problems))
    if not args.dry_run:
        dr.append(path, records)
    counts = {choice: sum(r["decision"] == choice for r in records) for choice in dr.REVIEW_DECISIONS}
    print(f"{'Would save' if args.dry_run else 'Saved'} {len(records)} decisions from {args.sheet} "
          f"({', '.join(f'{k} {v}' for k, v in counts.items())}); {len(rows) - len(records)} rows left blank")
    return 0


def usable_records(db, records: list[dict]) -> tuple[list[dict], int]:
    """Drop reviews of Zomi that has changed since it was shown; their approvals then have nothing to rest on."""
    digests = current_digests(db, sorted({r["sense_id"] for r in records}))
    kept, stale = [], 0
    for record in records:
        if dr.kind_of(record) == "review" and digests.get(record["sense_id"]) != record["shown"]:
            stale += 1
            continue
        kept.append(record)
    return kept, stale


def cmd_approval_sheet(args: argparse.Namespace) -> int:
    from sqlalchemy import text

    records = dr.load(args.reviews_dir / "dictionary-reviews.jsonl")
    with session() as db:
        records, stale = usable_records(db, records)
        decisions = dr.decide(records, args.min_reviewers)
        ready = sorted(sense for sense, d in decisions.items() if d.status == dr.READY_FOR_APPROVAL)
        rows = []
        if ready:
            found = {row.id: row for row in db.execute(text(SENSE_SQL + " WHERE s.id = ANY(:ids)"), {"ids": ready})}
            for sense in ready:
                row = found.get(sense)
                if row is None:
                    continue
                rows.append([sense, row.headword, row.part_of_speech or "", row.definition_en or "", zomi_of(row),
                             ", ".join(decisions[sense].reviewers), None, ""])
    if not rows:
        print("No senses are ready for approval.")
        return 0
    batches = load_batches(args.reviews_dir / "dictionary-batches.jsonl")
    name = args.batch or next_batch_name(batches, "dict-approval")
    about = {"kind": "dictionary-approval", "batch": name, "exported_at": now_utc()}
    write_sheet(args.out, "Approve", APPROVE_COLUMNS, rows, dr.APPROVAL_DECISIONS, about, "Decision",
                ("Decision", "Notes"))
    dr.append(args.reviews_dir / "dictionary-batches.jsonl", [{
        "kind": "approval-batch", "batch": name, "exported_at": about["exported_at"], "sense_ids": [r[0] for r in rows]}])
    print(f"Wrote {len(rows)} senses ready for approval to {args.out} (batch {name}); {stale} stale reviews ignored")
    return 0


def cmd_approve(args: argparse.Namespace) -> int:
    rows, about = read_sheet(args.sheet, "Approve")
    if about.get("kind") != "dictionary-approval":
        sys.exit(f"{args.sheet} is not a dictionary approval sheet")
    path = args.reviews_dir / "dictionary-reviews.jsonl"
    roster = dr.load_roster(args.reviews_dir / "roster.csv")
    reviewers = {sense: set(d.reviewers) for sense, d in dr.decide(dr.load(path)).items()}
    records, problems = [], []
    for number, row in enumerate(rows, start=2):
        decision = row.get("Decision", "").strip().lower()
        if not decision:
            continue
        sense = row.get("Sense ID", "")
        record = {"kind": "approval", "sense_id": sense, "decision": decision, "approver": args.approver,
                  "approved_on": args.date, "recorded_at": now_utc(), "batch": about.get("batch", ""),
                  "notes": row.get("Notes", "")}
        errors = dr.validate(record)
        if args.approver in reviewers.get(sense, set()):
            errors.append(f"{args.approver} reviewed this sense and cannot also approve it")
        if errors:
            problems.append(f"row {number}: {'; '.join(errors)}")
        records.append(record)
    problems += dr.roster_problems(records, roster)
    if problems:
        sys.exit("Nothing was saved:\n  " + "\n  ".join(problems))
    if not args.dry_run:
        dr.append(path, records)
    print(f"{'Would save' if args.dry_run else 'Saved'} {len(records)} approval decisions from {args.sheet}")
    return 0


def cmd_apply(args: argparse.Namespace) -> int:
    from sqlalchemy import text

    records = dr.load(args.reviews_dir / "dictionary-reviews.jsonl")
    with session() as db:
        records, stale = usable_records(db, records)
        decisions = dr.decide(records, args.min_reviewers)
        counts: dict[str, int] = {}
        for sense, decision in decisions.items():
            review_status = dr.REVIEW_STATUS[decision.status]
            status = {"approved": "approved", "rejected": "rejected"}.get(review_status, "in_review")
            db.execute(text("UPDATE dictionary_senses SET review_status = :review, status = :status WHERE id = :id"),
                       {"review": review_status, "status": status, "id": sense})
            counts[decision.status] = counts.get(decision.status, 0) + 1
        if args.dry_run:
            db.rollback()
        else:
            db.commit()
    summary = ", ".join(f"{status} {counts.get(status, 0)}" for status in dr.STATUSES)
    print(f"{'Would set' if args.dry_run else 'Set'} {sum(counts.values())} senses: {summary}; "
          f"{stale} reviews of since-changed Zomi ignored")
    return 0


def cmd_status(args: argparse.Namespace) -> int:
    from sqlalchemy import text

    with session() as db:
        rows = db.execute(text("SELECT review_status, count(*) FROM dictionary_senses GROUP BY 1 ORDER BY 1")).all()
    for review_status, count in rows:
        print(f"{review_status:18} {count:,}")
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--reviews-dir", type=Path, default=None,
                        help="private folder with roster.csv and the decisions files (default: ../siamsil-data/reviews)")
    commands = parser.add_subparsers(dest="command", required=True)

    export = commands.add_parser("export", help="a review sheet for one reviewer")
    export.add_argument("--reviewer", required=True, help="roster code, e.g. R01")
    export.add_argument("--out", type=Path, required=True)
    export.add_argument("--size", type=int, default=500)
    export.add_argument("--source", help="only senses from this source id (default: the primary source)")
    export.add_argument("--all-sources", action="store_true", help="senses from every source, not only the primary")
    export.add_argument("--issue", help="only senses with this open issue code, e.g. VARIETY_CHECK")
    export.add_argument("--batch", help="batch name (default: dict-<reviewer>-NNNN)")
    export.add_argument("--again", action="store_true", help="include senses already sent in earlier batches")

    record = commands.add_parser("record", help="save a completed review sheet")
    record.add_argument("--sheet", type=Path, required=True)
    record.add_argument("--reviewer", required=True)
    record.add_argument("--date", default=date.today().isoformat(), help="day the reviews were made")
    record.add_argument("--dry-run", action="store_true")

    sheet = commands.add_parser("approval-sheet", help="senses ready for a lead editor")
    sheet.add_argument("--out", type=Path, required=True)
    sheet.add_argument("--batch")
    sheet.add_argument("--min-reviewers", type=int, default=1)

    approve = commands.add_parser("approve", help="save a lead editor's decisions")
    approve.add_argument("--sheet", type=Path, required=True)
    approve.add_argument("--approver", required=True, help="roster code, e.g. E01")
    approve.add_argument("--date", default=date.today().isoformat())
    approve.add_argument("--dry-run", action="store_true")

    apply = commands.add_parser("apply", help="set review statuses in the database")
    apply.add_argument("--min-reviewers", type=int, default=1)
    apply.add_argument("--dry-run", action="store_true")

    commands.add_parser("status", help="senses per review status")

    args = parser.parse_args(argv)
    args.reviews_dir = args.reviews_dir or default_reviews_dir()
    handlers = {"export": cmd_export, "record": cmd_record, "approval-sheet": cmd_approval_sheet,
                "approve": cmd_approve, "apply": cmd_apply, "status": cmd_status}
    try:
        return handlers[args.command](args)
    except ValueError as error:
        sys.exit(f"Nothing was changed: {error}")


if __name__ == "__main__":
    raise SystemExit(main())
