#!/usr/bin/env python3
"""Count how Zolai sources actually spell things, as evidence for the Siamsil Zolai Standard.

Reads Zomi text from several sources, counts every word form per source, and reports:

  forms.csv            every word form with its count in each source
  variant_families.csv forms that differ only by apostrophes, hyphens, or doubled vowels
  joins.csv            words written both joined and split ("ab" / "a b" / "a-b")
  rule_checks.csv      the spellings each style-guide point is about (data/zolai_standard/rule_checks.json)
  Siamsil_Zolai_Standard_Review.xlsx  the same, with decision columns for the language reviewers
  summary.md           sources, sizes, and the most divided points

Nothing is corrected or decided here: the reports only show what the sources do, so Zomi
reviewers can decide (docs/ZOLAI_STANDARD.md). Source text is copyrighted and private;
write the output to the private siamsil-data repository, not this one.

  python3 ml_pipeline/scripts/spelling_variants.py --out ../siamsil-data/standard/derived \\
      --bible tdb=data/processed/bible/bible_versions.sqlite:tdb \\
      --text zokam=../siamsil-data/books/derived/zokam-le-zolai.txt \\
      --text primer=../siamsil-data/grammar/Paunam_Khenna_Leh_Kampau_Luanzia.pdf \\
      --xlsx dictionary="../siamsil-data/dictionary/Zomi_Standard_Dictionary_AI_Cleaned.xlsx:Definition (Zomi)"

PDF text is read with Poppler's pdftotext. Order the sources as they should appear in the reports.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sqlite3
import subprocess
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Iterator

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_CHECKS = ROOT / "data" / "zolai_standard" / "rule_checks.json"

APOSTROPHES = str.maketrans({"’": "'", "‘": "'", "ʼ": "'", "`": "'", "´": "'"})
DASHES = str.maketrans({"‐": "-", "‑": "-"})
# A word: letters, with inner apostrophes or hyphens, and an optional final apostrophe (possessive).
TOKEN = re.compile(r"[a-z]+(?:['-][a-z]+)*'?")
DOUBLED_VOWEL = re.compile(r"([aeiou])\1+")


def tokens(text: str) -> list[str]:
    return TOKEN.findall(text.lower().translate(APOSTROPHES).translate(DASHES))


def family_key(form: str) -> str:
    """Spellings that differ only by apostrophes, hyphens, or doubled vowels share a key."""
    return DOUBLED_VOWEL.sub(r"\1", form.replace("'", "").replace("-", ""))


@dataclass
class Source:
    name: str
    description: str
    words: Counter = field(default_factory=Counter)
    pairs: Counter = field(default_factory=Counter)
    ngrams: Counter = field(default_factory=Counter)  # 3-grams, for three-word variants
    possessive_apostrophes: int = 0

    def add(self, text: str) -> None:
        words = tokens(text)
        self.words.update(words)
        self.pairs.update(zip(words, words[1:]))
        self.ngrams.update(zip(words, words[1:], words[2:]))
        self.possessive_apostrophes += sum(1 for word in words if word.endswith("'"))

    def count(self, phrase: str) -> int:
        words = tokens(phrase)
        if len(words) == 1:
            return self.words[words[0]]
        if len(words) == 2:
            return self.pairs[tuple(words)]
        if len(words) == 3:
            return self.ngrams[tuple(words)]
        raise ValueError(f"variants are one to three words: {phrase!r}")


# Reading sources ---------------------------------------------------------------------------------


def read_text(path: Path) -> str:
    if path.suffix.lower() == ".pdf":
        try:
            result = subprocess.run(["pdftotext", "-enc", "UTF-8", str(path), "-"], capture_output=True, check=True)
        except FileNotFoundError:
            sys.exit("pdftotext (Poppler) is needed to read PDF sources")
        return result.stdout.decode("utf-8", errors="replace")
    return path.read_text(encoding="utf-8")


def bible_verses(path: Path, version: str) -> Iterator[str]:
    con = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
    try:
        rows = con.execute("SELECT text FROM verses WHERE version_id = ? ORDER BY book_id, chapter, verse", (version,))
        yield from (row[0] for row in rows)
    finally:
        con.close()


def xlsx_column(path: Path, column: str) -> Iterator[str]:
    import openpyxl

    sheet = openpyxl.load_workbook(path, read_only=True).worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = [str(cell or "").strip() for cell in next(rows)]
    if column not in header:
        sys.exit(f"{path.name} has no column {column!r}; columns: {', '.join(header)}")
    index = header.index(column)
    for row in rows:
        if row and row[index]:
            yield str(row[index])


def split_spec(spec: str, flag: str) -> tuple[str, str]:
    name, sep, rest = spec.partition("=")
    if not sep or not name or not rest:
        sys.exit(f"{flag} expects NAME=VALUE, got {spec!r}")
    return name, rest


def load_sources(args: argparse.Namespace) -> list[Source]:
    sources: list[Source] = []
    for flag, spec in args.ordered:
        name, value = split_spec(spec, flag)
        if any(source.name == name for source in sources):
            sys.exit(f"source name {name!r} is used twice")
        if flag == "--text":
            path = Path(value)
            source = Source(name, path.name)
            source.add(read_text(path))
        elif flag == "--bible":
            path, _, version = value.rpartition(":")
            source = Source(name, f"{Path(path).name}, version {version}")
            for verse in bible_verses(Path(path), version):
                source.add(verse)
        else:
            path, _, column = value.rpartition(":")
            source = Source(name, f"{Path(path).name}, column {column}")
            for cell in xlsx_column(Path(path), column):
                source.add(cell)
        if not source.words:
            sys.exit(f"source {name!r} has no words")
        sources.append(source)
    return sources


# Analyses ----------------------------------------------------------------------------------------


def all_words(sources: list[Source]) -> Counter:
    total: Counter = Counter()
    for source in sources:
        total.update(source.words)
    return total


def base_counts(words: Counter) -> Counter:
    """Counts with a final possessive apostrophe folded in: "topa'" counts as "topa"."""
    folded: Counter = Counter()
    for form, count in words.items():
        folded[form.rstrip("'")] += count
    return folded


def variant_families(sources: list[Source], min_total: int) -> list[dict]:
    """Groups of two or more spellings sharing a family key, most divided groups first. A final
    apostrophe is grammar (the possessive), not a spelling difference, so it is folded in first."""
    groups: dict[str, set[str]] = defaultdict(set)
    totals = base_counts(all_words(sources))
    for form in totals:
        groups[family_key(form)].add(form)
    families = []
    for key, forms in groups.items():
        if len(forms) < 2:
            continue
        total = sum(totals[form] for form in forms)
        if total < min_total:
            continue
        ranked = sorted(forms, key=lambda form: (-totals[form], form))
        families.append({"key": key, "forms": ranked, "total": total,
                         "minority": total - totals[ranked[0]]})
    return sorted(families, key=lambda family: (-family["minority"], -family["total"], family["key"]))


def joins(sources: list[Source], min_total: int) -> list[dict]:
    """Words seen both joined ("ab") and as two words ("a b") or hyphenated ("a-b")."""
    words = all_words(sources)
    pairs: Counter = Counter()
    for source in sources:
        pairs.update(source.pairs)
    rows = []
    for joined in words:
        if "'" in joined or "-" in joined or len(joined) < 3:
            continue
        best = None
        for cut in range(1, len(joined)):
            left, right = joined[:cut], joined[cut:]
            split = pairs[(left, right)]
            hyphen = words[f"{left}-{right}"]
            if split + hyphen and (best is None or split + hyphen > best[2] + best[3]):
                best = (left, right, split, hyphen)
        if best is None:
            continue
        left, right, split, hyphen = best
        total = words[joined] + split + hyphen
        if total < min_total:
            continue
        rows.append({"joined": joined, "split": f"{left} {right}", "hyphenated": f"{left}-{right}",
                     "total": total, "minority": total - max(words[joined], split, hyphen)})
    return sorted(rows, key=lambda row: (-row["minority"], -row["total"], row["joined"]))


def rule_checks(sources: list[Source], checks: list[dict]) -> list[dict]:
    rows = []
    for check in checks:
        for number, variants in enumerate(check["variants"], start=1):
            for variant in variants:
                counts = {}
                for source in sources:
                    if check.get("special") == "possessive" and variant == "'":
                        counts[source.name] = source.possessive_apostrophes
                    else:
                        counts[source.name] = source.count(variant)
                label = "word ending in an apostrophe" if variant == "'" else variant
                rows.append({"check": check["id"], "set": number, "variant": label, "counts": counts})
    return rows


# Suggestions -------------------------------------------------------------------------------------

SUGGEST_MIN_USES = 3  # a source with fewer uses of a word does not vote
SPLIT_SHARE = 0.2  # a source using a second spelling this often may mean two different words
SPLIT_MIN_USES = 20  # ... judged only where the source uses the word this often, so a handful of uses cannot block


def suggest(counts: dict[str, dict[str, int]], forms: list[str], reference: str) -> tuple[str, str, str]:
    """Siamsil's suggestion for a reviewer, from the counts alone: (spelling, strength, why).

    Each source with enough uses votes for the spelling it uses most; the spelling most sources
    prefer is suggested, and the reference source breaks ties. When one source itself uses two
    spellings often, they may be different words, so nothing is suggested. A draft for the
    reviewers, never a decision.
    """
    preferred: dict[str, str] = {}
    for source, by_form in counts.items():
        uses = sum(by_form.get(form, 0) for form in forms)
        if uses < SUGGEST_MIN_USES:
            continue
        ranked = sorted(forms, key=lambda form: -by_form.get(form, 0))
        if len(ranked) > 1 and uses >= SPLIT_MIN_USES and by_form.get(ranked[1], 0) / uses >= SPLIT_SHARE:
            return "", "check meaning", (f"{source} uses both {ranked[0]} ({by_form[ranked[0]]}) and "
                                         f"{ranked[1]} ({by_form[ranked[1]]}): different words, or {source} is inconsistent")
        preferred[source] = ranked[0]
    if not preferred:
        return "", "too few uses", f"no source uses it {SUGGEST_MIN_USES} or more times"
    votes = Counter(preferred.values())
    best = max(votes, key=lambda form: (votes[form], form == preferred.get(reference)))
    agree = [source for source, form in preferred.items() if form == best]
    others = [f"{source} ({form})" for source, form in preferred.items() if form != best]
    why = f"preferred by {', '.join(agree)}" + (f"; other spellings in {', '.join(others)}" if others else "")
    return best, "strong" if not others else "mixed", why


# Output ------------------------------------------------------------------------------------------


def per_source(form_or_phrase: str, sources: list[Source]) -> list[int]:
    return [source.count(form_or_phrase) for source in sources]


def write_csv(path: Path, header: list[str], rows: Iterable[list]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        writer.writerows(rows)


def write_reports(out: Path, sources: list[Source], checks: list[dict], min_total: int, reference: str | None = None) -> dict:
    out.mkdir(parents=True, exist_ok=True)
    names = [source.name for source in sources]
    reference = reference or names[0]
    if reference not in names:
        raise ValueError(f"reference source {reference!r} is not one of {', '.join(names)}")
    totals = all_words(sources)

    write_csv(out / "forms.csv", ["form", "total", *names],
              ([form, totals[form], *(source.words[form] for source in sources)]
               for form in sorted(totals, key=lambda form: (-totals[form], form))))

    families = variant_families(sources, min_total)
    folded = [base_counts(source.words) for source in sources]
    write_csv(out / "variant_families.csv", ["family", "form", "total", *names],
              ([family["key"], form, sum(counts[form] for counts in folded), *(counts[form] for counts in folded)]
               for family in families for form in family["forms"]))

    joined = joins(sources, min_total)
    write_csv(out / "joins.csv", ["joined", "spelling", "form", "total", *names],
              ([row["joined"], kind, row[kind], sum(per_source(row[kind], sources)), *per_source(row[kind], sources)]
               for row in joined for kind in ("joined", "split", "hyphenated")))

    checked = rule_checks(sources, checks)
    write_csv(out / "rule_checks.csv", ["check", "set", "variant", "total", *names],
              ([row["check"], row["set"], row["variant"], sum(row["counts"].values()),
                *(row["counts"][name] for name in names)] for row in checked))

    write_workbook(out / "Siamsil_Zolai_Standard_Review.xlsx", sources, checks, checked, joined, families, reference)

    summary = {
        "sources": [{"name": source.name, "description": source.description, "words": sum(source.words.values()),
                     "distinct_forms": len(source.words)} for source in sources],
        "distinct_forms": len(totals),
        "variant_families": len(families),
        "divided_joins": sum(1 for row in joined if row["minority"]),
        "rule_checks": len(checks),
        "min_total": min_total,
    }
    (out / "summary.json").write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    write_summary(out / "summary.md", summary, sources, checks, checked, joined, families)
    return summary


def write_summary(path: Path, summary: dict, sources, checks, checked, joined, families) -> None:
    names = [source.name for source in sources]
    lines = ["# Zolai spelling evidence", "",
             "Made by `ml_pipeline/scripts/spelling_variants.py`. Counts only: nothing here is a decision.", "",
             "## Sources", "", "| Source | Read from | Words | Distinct forms |", "| --- | --- | ---: | ---: |"]
    lines += [f"| {s['name']} | {s['description']} | {s['words']:,} | {s['distinct_forms']:,} |" for s in summary["sources"]]
    lines += ["", "## Style-guide points", "", f"| Check | Variant | {' | '.join(names)} |",
              f"| --- | --- | {' | '.join('---:' for _ in names)} |"]
    lines += [f"| {row['check']} | {row['variant']} | {' | '.join(str(row['counts'][n]) for n in names)} |" for row in checked]
    lines += ["", f"## Most divided joined/split spellings (of {summary['divided_joins']:,})", "",
              "| Joined | Split | Hyphenated | Total |", "| --- | ---: | ---: | ---: |"]
    words = all_words(sources)
    pairs: Counter = Counter()
    for source in sources:
        pairs.update(source.pairs)
    for row in joined[:40]:
        left, right = row["split"].split(" ")
        lines.append(f"| {row['joined']} ({words[row['joined']]}) | {row['split']} ({pairs[(left, right)]}) "
                     f"| {row['hyphenated']} ({words[row['hyphenated']]}) | {row['total']} |")
    lines += ["", f"## Most divided spelling families (of {summary['variant_families']:,})", "",
              "Forms that differ only by inner apostrophes, hyphens, or doubled vowels (a final possessive apostrophe is counted with its word). Some are different words "
              "(a long vowel can change the meaning); reviewers decide which are variants.", "",
              "| Forms (count) | Total |", "| --- | ---: |"]
    folded = base_counts(words)
    for family in families[:40]:
        lines.append(f"| {', '.join(f'{form} ({folded[form]})' for form in family['forms'])} | {family['total']} |")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_workbook(path: Path, sources, checks, checked, joined, families, reference: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    names = [source.name for source in sources]
    words = all_words(sources)
    bold = Font(bold=True)
    head = PatternFill("solid", fgColor="1F3A5F")
    white = Font(bold=True, color="FFFFFF")
    decide = PatternFill("solid", fgColor="FFF4CC")
    wrap = Alignment(wrap_text=True, vertical="top")

    book = Workbook()
    intro = book.active
    intro.title = "Read me"
    for line in (
        "Siamsil Zolai Standard: spelling evidence for review",
        "",
        "Each row shows how the sources spell something, and what the three style guides say.",
        "Fill in the yellow columns: the spelling Siamsil should use, and why. Leave a row blank if unsure.",
        "Use your reviewer code (R01, E01), never your name. Two people must agree before a rule is adopted.",
        "The 'Siamsil suggestion (AI)' columns are drafts to save time, not decisions: on Rules they weigh the guides and",
        f"the counts; on Joins and Families they are the spelling most sources prefer ({reference} breaks ties). Overrule freely.",
        "",
        "Rules: the points the guides address. Joins: words written both joined and split. Families: spellings",
        "differing only by apostrophes, hyphens, or doubled vowels (some are different words).",
        f"Joins and Families show the {WORKBOOK_ROWS} most divided; joins.csv and variant_families.csv have them all.",
        "",
        "Sources: " + "; ".join(f"{s.name} = {s.description}" for s in sources),
    ):
        intro.append([line])
    intro["A1"].font = Font(bold=True, size=14)
    intro.column_dimensions["A"].width = 120

    def table(title: str, header: list[str], rows: list[list], widths: dict[str, int], choices: list[str] | None):
        sheet = book.create_sheet(title)
        sheet.append(header)
        for cell in sheet[1]:
            cell.font, cell.fill, cell.alignment = white, head, wrap
        for row in rows:
            sheet.append(row)
        decision_columns = [index for index, name in enumerate(header, start=1) if name in DECISION_HEADERS]
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = wrap
            for index in decision_columns:
                row[index - 1].fill = decide
        for letter, width in widths.items():
            sheet.column_dimensions[letter].width = width
        sheet.freeze_panes = "A2"
        if choices and sheet.max_row > 1:
            column = header.index("Decision") + 1
            validation = DataValidation(type="list", formula1=f'"{",".join(choices)}"', allow_blank=True)
            sheet.add_data_validation(validation)
            letter = sheet.cell(row=1, column=column).column_letter
            validation.add(f"{letter}2:{letter}{sheet.max_row}")
        return sheet

    by_check = {check["id"]: check for check in checks}
    rule_rows = []
    sets: dict[tuple[str, int], list[dict]] = defaultdict(list)
    for row in checked:
        sets[(row["check"], row["set"])].append(row)
    for (check_id, _), variants in sets.items():
        check = by_check[check_id]
        counts = "\n".join(f"{v['variant']}: " + ", ".join(f"{n} {v['counts'][n]}" for n in names) for v in variants)
        suggestion = check.get("suggestion", {})
        rule_rows.append([check_id, check["topic"], " / ".join(v["variant"] for v in variants),
                          check["zsf"], check["zokam"], check["primer"], counts,
                          "\n".join(text for text in (check.get("note"), check.get("owner_input")) if text),
                          suggestion.get("spelling", ""), suggestion.get("reason", ""), "", "", "", ""])
    table("Rules", ["Check", "Topic", "Spellings", "Zolai Standard Format", "Zokam le Zolai", "Primer",
                    "Counts in the sources", "Note", "Siamsil suggestion (AI)", "Why",
                    "Decision", "Spelling to use", "Reason", "Reviewer"],
          rule_rows, {"A": 16, "B": 30, "C": 22, "D": 40, "E": 40, "F": 40, "G": 46, "H": 30, "I": 24, "J": 44,
                      "K": 14, "L": 18, "M": 30, "N": 10},
          ["Adopt", "Allow both", "Need more evidence"])

    pairs: Counter = Counter()
    for source in sources:
        pairs.update(source.pairs)
    join_rows = []
    for row in joined[: WORKBOOK_ROWS * 2]:
        if not row["minority"]:
            continue
        left, right = row["split"].split(" ")
        spellings = [row["joined"], row["split"], row["hyphenated"]]
        counts = {source.name: {row["joined"]: source.words[row["joined"]], row["split"]: source.pairs[(left, right)],
                                row["hyphenated"]: source.words[row["hyphenated"]]} for source in sources}
        best, strength, why = suggest(counts, spellings, reference)
        join_rows.append([row["joined"], words[row["joined"]], pairs[(left, right)], words[row["hyphenated"]],
                          row["split"], row["hyphenated"], best, strength, why, "", "", ""])
    table("Joins", ["Joined", "Joined count", "Split count", "Hyphen count", "Split form", "Hyphenated form",
                    "Siamsil suggestion (AI)", "Strength", "Why", "Decision", "Reason", "Reviewer"],
          join_rows[:WORKBOOK_ROWS], {"A": 22, "E": 22, "F": 22, "G": 22, "H": 14, "I": 44, "J": 14, "K": 30},
          ["Joined", "Split", "Hyphenated", "Depends on meaning"])

    folded = base_counts(words)
    per_source_folded = {source.name: base_counts(source.words) for source in sources}
    family_rows = []
    for family in families[:WORKBOOK_ROWS]:
        best, strength, why = suggest(per_source_folded, family["forms"], reference)
        family_rows.append([", ".join(f"{form} ({folded[form]})" for form in family["forms"]), family["total"],
                            best, strength, why, "", "", ""])
    table("Families", ["Forms (count)", "Total", "Siamsil suggestion (AI)", "Strength", "Why",
                       "Decision", "Spelling to use", "Reviewer"], family_rows,
          {"A": 60, "C": 20, "D": 14, "E": 50, "F": 16, "G": 18}, ["Same word", "Different words"])

    sources_sheet = book.create_sheet("Sources")
    sources_sheet.append(["Name", "Read from", "Words", "Distinct forms"])
    for cell in sources_sheet[1]:
        cell.font = bold
    for source in sources:
        sources_sheet.append([source.name, source.description, sum(source.words.values()), len(source.words)])
    book.save(path)


DECISION_HEADERS = {"Decision", "Spelling to use", "Reason", "Reviewer"}  # filled in by reviewers
WORKBOOK_ROWS = 500  # rows per sheet a reviewer can work through; the CSV files have everything


def main() -> int:
    parser = argparse.ArgumentParser(description="Count how Zolai sources spell things (evidence for the Siamsil Zolai Standard).")
    parser.add_argument("--text", action="append", default=[], metavar="NAME=PATH", help="a .txt or .pdf source")
    parser.add_argument("--bible", action="append", default=[], metavar="NAME=SQLITE:VERSION",
                        help="a version in bible_versions.sqlite (import_usx_bibles.py)")
    parser.add_argument("--xlsx", action="append", default=[], metavar="NAME=PATH:COLUMN", help="a text column of a workbook")
    parser.add_argument("--checks", type=Path, default=DEFAULT_CHECKS, help="style-guide points to count")
    parser.add_argument("--min-total", type=int, default=5, help="ignore joins and families seen fewer times")
    parser.add_argument("--reference", help="source that breaks ties in the suggestions (default: the first)")
    parser.add_argument("--out", type=Path, required=True, help="output folder (private: siamsil-data)")
    args = parser.parse_args()

    # Keep the sources in command-line order.
    args.ordered = []
    argv = sys.argv[1:]
    for index, item in enumerate(argv):
        for flag in ("--text", "--bible", "--xlsx"):
            if item == flag and index + 1 < len(argv):
                args.ordered.append((flag, argv[index + 1]))
            elif item.startswith(flag + "="):
                args.ordered.append((flag, item[len(flag) + 1:]))
    if not args.ordered:
        parser.error("give at least one source with --text, --bible, or --xlsx")

    checks = json.loads(args.checks.read_text(encoding="utf-8"))["checks"]
    sources = load_sources(args)
    try:
        summary = write_reports(args.out, sources, checks, args.min_total, args.reference)
    except ValueError as error:
        sys.exit(str(error))
    for source in summary["sources"]:
        print(f"{source['name']:<12} {source['words']:>10,} words  {source['distinct_forms']:>8,} forms")
    print(f"{summary['variant_families']:,} spelling families, {summary['divided_joins']:,} divided joins, "
          f"{summary['rule_checks']} style-guide points -> {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
